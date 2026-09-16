#!/usr/bin/env python3
"""Deterministic SOF hanger validation and workbook update utility.

The tool deliberately treats LLM output as untrusted rule input. Rules only become
eligible after their schema, required values, selected source workbook, source
sheet, and source cell range have been validated.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries


ORDER_FIELDS = [
    "Division", "Season", "Year", "Ref#", "Style", "Product Description",
    "Color", "Label", "Label Name", "Account", "PO#", "PO Qty",
    "Size Configuration", "Pack Ratio", "Master Box Quantity", "Hang/Flat",
    "Hanger code", "Hanger color", "Color Sizer", "Sticker hanger",
    "SIZE Sticker hanger",
]

ALIASES = {
    "division": "Division", "season": "Season", "year": "Year",
    "ref#": "Ref#", "ref": "Ref#", "ref number": "Ref#",
    "style": "Style", "product description": "Product Description",
    "color": "Color", "label": "Label", "label name": "Label Name",
    "account": "Account",
    "po#": "PO#", "po": "PO#", "po number": "PO#", "po qty": "PO Qty",
    "size configuration": "Size Configuration", "pack ratio": "Pack Ratio",
    "master box quantity": "Master Box Quantity", "hang/flat": "Hang/Flat",
    "hanger code": "Hanger code", "hanger color": "Hanger color",
    "color sizer": "Color Sizer", "sticker hanger": "Sticker hanger",
    "size sticker hanger": "SIZE Sticker hanger",
}

# The only sheets a hanger decision may be sourced from. Sheets such as
# "General Info" and "REPLESNISHMENT" carry general packing rules that are
# qualified by phrases like "except the categories below", so a citation from
# them is not evidence about one style and is rejected.
CATEGORY_SOURCE_SHEETS = {
    "HUGGIES", "SWEATER-YARN", "TOPS", "BOTTOMS", "COVERALLS", "SETS",
    "OUTERWEAR", "GIRLS SWIMWEAR", "BOX SETS", "LACOSTE TOYS", "HOSIERY",
    "UNDERWEAR", "COLD WEATHER", "BAGS", "BODYSUITS", "DRESSES",
}

# Some account SOFs use a label-driven packing table instead of one sheet per
# product category. PACKAGING is accepted for LLM decisions only when the cited
# cells contain the order's exact Label and a direct Flat/Hanger statement.
LLM_SOURCE_SHEETS = CATEGORY_SOURCE_SHEETS | {"PACKAGING"}

WRITE_FIELDS = [
    "Hanger code", "Hanger color", "Color Sizer", "Sticker hanger",
    "SIZE Sticker hanger",
]

# Colour convention used by "TONG HOP HANGER" workbooks: order-file columns are
# cyan, columns sourced from the SOF are yellow. Only the header row is filled.
SOF_HEADER_FILL = "FFFF00"
ORDER_HEADER_FILL = "33CCCC"

# Single-sheet deliverable, laid out to match the TONG HOP HANGER workbooks so it
# can be pasted straight in: order columns first, then the SOF block.
RESULT_ORDER_COLUMNS = [
    "Division", "Season", "Year", "Ref#", "Style", "Product Description",
    "Color", "Label", "Label Name", "Account", "PO#", "PO Qty",
    "Size Configuration", "Pack Ratio", "Master Box Quantity", "Hang/Flat",
]
# NCC is filled in by hand: it names the supplier and appears nowhere in the SOF,
# so the column is created and coloured but deliberately left empty.
RESULT_SOF_COLUMNS = [
    "Hanger code", "Hanger color", "Color Sizer", "Sticker hanger",
    "SIZE Sticker hanger", "NCC",
]
RESULT_SOF_SOURCE = {
    "Hanger code": "hanger_code",
    "Hanger color": "hanger_color",
    "Color Sizer": "color_sizer",
    "Sticker hanger": "sticker_hanger",
    "SIZE Sticker hanger": "size_sticker_hanger",
}

# Traceability block. Kept grey rather than cyan/yellow so it reads as "not part
# of the TONG HOP paste block" when someone copies columns into that workbook.
AUDIT_HEADER_FILL = "D9D9D9"
RESULT_AUDIT_COLUMNS = [
    ("Status", "status"),
    ("Validation Note", "validation_note"),
    ("Match Method", "match_method"),
    ("Confidence", "confidence"),
    ("SOF File", "source_file"),
    ("SOF Sheet", "source_sheet"),
    ("SOF Cells", "source_cells"),
    ("SOF Section", "source_section"),
    ("SOF Lines", "source_lines"),
    ("SOF Excerpt", "source_excerpt"),
]

REVIEW_HEADERS = [
    "Order Sheet", "Order Row", "PO#", "Account", "Division", "Label",
    "Label Name", "Ref#", "Style", "Product Description", "Product Category",
    "Size Configuration", "Order Hang/Flat", "SOF Hang/Flat",
    "Hanger code", "Hanger color", "Color Sizer", "Sticker hanger",
    "SIZE Sticker hanger", "Validation Status", "Validation Note",
    "SOF Source", "Match Method", "Confidence",
]

VALID_HANG_FLAT = {"Hang", "Flat"}
SOURCE_RANGE_RE = re.compile(
    r"^(?:\$?[A-Z]{1,3}\$?[1-9]\d*)(?::(?:\$?[A-Z]{1,3}\$?[1-9]\d*))?$"
)

# A SOF states the hanger code in a table but the hanger colour in a sentence
# below it, so a decision may cite several ranges. The cap keeps a citation
# reviewable by hand instead of letting a whole sheet be quoted as "evidence".
MAX_SOURCE_RANGES = 3

# Value reserved for the LLM to say a field is absent from the SOF. It is never
# written to a workbook.
NOT_STATED = "NOT_STATED"

# A SOF names an accessory only where it is required, so for these three
# "not stated" and "none required" are the same fact and NOT_STATED collapses to
# NO. For the hanger itself they are not the same: an unstated code or colour
# means the governing rule was never found, so it stays NOT_STATED and the row
# is reviewed.
ACCESSORY_FIELDS = ("color_sizer", "sticker_hanger", "size_sticker_hanger")


class AutomationError(RuntimeError):
    pass


def clean(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def key(value: Any) -> str:
    return clean(value).casefold()


def normalize_size(value: Any) -> str:
    text = key(value)
    text = re.sub(r"\s*[-–—/]\s*", "-", text)
    text = re.sub(r"\s+", "", text)
    return text


def normalize_hang_flat(value: Any) -> str | None:
    text = key(value)
    if text in {"hang", "hanger", "hung"}:
        return "Hang"
    if text in {"flat", "flat packed", "flat-packed", "flat pack"}:
        return "Flat"
    return None


def displayed_identifier(cell: Cell) -> str:
    """Return an identifier as displayed, including simple leading-zero formats."""
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        fmt = clean(cell.number_format)
        if re.fullmatch(r"0+", fmt):
            return f"{value:0{len(fmt)}d}"
    if isinstance(value, float) and value.is_integer():
        fmt = clean(cell.number_format)
        if re.fullmatch(r"0+", fmt):
            return f"{int(value):0{len(fmt)}d}"
        return str(int(value))
    return clean(value)


def canonical_header(value: Any) -> str | None:
    text = key(value).replace("\n", " ")
    text = re.sub(r"\s+", " ", text)
    return ALIASES.get(text)


def find_header(ws, scan_rows: int = 50) -> tuple[int, dict[str, int]]:
    best: tuple[int, int, dict[str, int]] | None = None
    for row in range(1, min(ws.max_row, scan_rows) + 1):
        mapping: dict[str, int] = {}
        for col in range(1, ws.max_column + 1):
            name = canonical_header(ws.cell(row, col).value)
            if name and name not in mapping:
                mapping[name] = col
        score = len(set(mapping) & set(ORDER_FIELDS))
        if best is None or score > best[0]:
            best = (score, row, mapping)
    if not best or best[0] < 5 or "Account" not in best[2] or "Hang/Flat" not in best[2]:
        raise AutomationError(f"Could not identify an order header in sheet {ws.title!r}")
    return best[1], best[2]


def select_order_sheet(wb, requested: str | None, scan_rows: int):
    if requested:
        if requested not in wb.sheetnames:
            raise AutomationError(f"Order sheet {requested!r} does not exist")
        ws = wb[requested]
        header_row, columns = find_header(ws, scan_rows)
        return ws, header_row, columns
    candidates = []
    for ws in wb.worksheets:
        if ws.title == "HANGER REVIEW":
            continue
        try:
            header_row, columns = find_header(ws, scan_rows)
            missing_write_fields = sum(field not in columns for field in WRITE_FIELDS)
            base_fields = len(set(columns) - set(WRITE_FIELDS))
            candidates.append((missing_write_fields, base_fields, ws, header_row, columns))
        except AutomationError:
            pass
    if not candidates:
        raise AutomationError("No worksheet contains a recognizable order table")
    # A raw/pending order sheet normally lacks all five output columns. Prefer it
    # over historical sheets that already contain completed hanger data.
    candidates.sort(key=lambda item: (item[0], item[1]), reverse=True)
    top_rank = candidates[0][:2]
    if len(candidates) > 1 and candidates[1][:2] == top_rank:
        names = ", ".join(item[2].title for item in candidates if item[:2] == top_rank)
        raise AutomationError(f"More than one order sheet is equally likely: {names}")
    _, _, ws, header_row, columns = candidates[0]
    return ws, header_row, columns


def parse_revision_date(filename: str) -> datetime | None:
    stem = Path(filename).stem
    patterns = [
        (r"(?<!\d)(20\d{2})[._-](\d{1,2})[._-](\d{1,2})(?!\d)", "%Y-%m-%d"),
        (r"(?<!\d)(\d{1,2})[._-](\d{1,2})[._-](20\d{2}|\d{2})(?!\d)", "%m-%d-%Y"),
    ]
    for pattern, order in patterns:
        matches = list(re.finditer(pattern, stem))
        if not matches:
            continue
        values = matches[-1].groups()
        try:
            if order == "%Y-%m-%d":
                year, month, day = map(int, values)
            else:
                month, day, year = map(int, values)
                if year < 100:
                    year += 2000
            return datetime(year, month, day)
        except ValueError:
            continue
    return None


@dataclass(frozen=True)
class SofSelection:
    status: str
    path: Path | None
    revision_date: str | None
    reason: str
    candidates: tuple[str, ...]


def find_sof(sof_root: Path, account: str) -> SofSelection:
    if not account:
        return SofSelection("REVIEW", None, None, "Account is blank", ())
    if not sof_root.exists():
        return SofSelection("REVIEW", None, None, f"SOF root is unavailable: {sof_root}", ())
    needle = account.casefold()
    matches = sorted(
        (p for p in sof_root.rglob("*") if p.is_file()
         and p.suffix.casefold() in {".xls", ".xlsx", ".docx", ".pdf"}
         and needle in p.name.casefold()),
        key=lambda p: str(p).casefold(),
    )
    names = tuple(str(p) for p in matches)
    if not matches:
        return SofSelection("REVIEW", None, None, f"No SOF file found for account {account}", names)
    if len(matches) == 1:
        revision = parse_revision_date(matches[0].name)
        return SofSelection("MATCHED", matches[0], revision.date().isoformat() if revision else None, "", names)
    dated = [(parse_revision_date(p.name), p) for p in matches]
    if any(date is None for date, _ in dated):
        return SofSelection("REVIEW", None, None,
                            "Several SOF files found and at least one revision date is not verifiable", names)
    latest = max(date for date, _ in dated if date is not None)
    winners = [p for date, p in dated if date == latest]
    if len(winners) != 1:
        return SofSelection("REVIEW", None, latest.date().isoformat(),
                            "Several SOF files share the latest revision date", names)
    return SofSelection("MATCHED", winners[0], latest.date().isoformat(), "", names)


def convert_xls_to_xlsx(source: Path, temp_dir: Path) -> Path:
    target = temp_dir / f"{source.stem}.xlsx"
    conversion_errors = []
    if os.name == "nt":
        excel = None
        try:
            import win32com.client  # type: ignore
            excel = win32com.client.DispatchEx("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False
            workbook = excel.Workbooks.Open(str(source.resolve()))
            try:
                workbook.SaveAs(str(target.resolve()), FileFormat=51)
            finally:
                workbook.Close(False)
            return target
        except ImportError:
            conversion_errors.append("pywin32 is not installed")
        except Exception as exc:
            conversion_errors.append(f"Microsoft Excel conversion failed: {exc}")
        finally:
            if excel is not None:
                try:
                    excel.Quit()
                except Exception:
                    pass
    office = shutil.which("libreoffice") or shutil.which("soffice")
    if office:
        result = subprocess.run(
            [office, "--headless", "--convert-to", "xlsx", "--outdir", str(temp_dir), str(source)],
            capture_output=True, text=True, timeout=180,
        )
        if result.returncode == 0 and target.exists():
            return target
        conversion_errors.append(f"LibreOffice conversion failed: {result.stderr.strip()}")
    raise AutomationError(
        ".xls input requires Microsoft Excel with pywin32 on Windows, or LibreOffice. "
        "Install one converter; the original file will remain unchanged. "
        + "; ".join(conversion_errors)
    )


def load_editable_workbook(path: Path, temp_dir: Path):
    editable = path
    if path.suffix.casefold() == ".xls":
        editable = convert_xls_to_xlsx(path, temp_dir)
    if editable.suffix.casefold() != ".xlsx":
        raise AutomationError("Order input must be .xls or .xlsx")
    return load_workbook(editable, data_only=False), editable


def load_sof_workbook(path: Path, temp_dir: Path):
    readable = convert_xls_to_xlsx(path, temp_dir) if path.suffix.casefold() == ".xls" else path
    return load_workbook(readable, data_only=True, read_only=False)


def row_values(ws, row: int, columns: dict[str, int]) -> dict[str, str]:
    result = {}
    for field in ORDER_FIELDS:
        col = columns.get(field)
        if not col:
            result[field] = ""
        elif field in {"PO#", "Account", "Ref#", "Style"}:
            result[field] = displayed_identifier(ws.cell(row, col))
        else:
            result[field] = clean(ws.cell(row, col).value)
    return result


def rule_priority(match: dict[str, Any]) -> tuple[int, str]:
    if match.get("ref_number") and (match.get("style") or match.get("style_any")):
        return 0, "REF_STYLE_OVERRIDE"
    if match.get("ref_number"):
        return 1, "REF_OVERRIDE"
    if match.get("style") or match.get("style_any"):
        return 2, "STYLE_OVERRIDE"
    has_full = all(match.get(name) is not None for name in
                   ("division", "label", "product_category", "size_configuration"))
    if has_full:
        return 3, "ACCOUNT_DIVISION_LABEL_CATEGORY_SIZE"
    has_no_label = all(match.get(name) is not None for name in
                       ("division", "product_category", "size_configuration"))
    if has_no_label:
        return 4, "ACCOUNT_DIVISION_CATEGORY_SIZE"
    if match.get("product_category"):
        return 5, "GENERAL_CATEGORY"
    return 99, "INVALID"


def field_matches(actual: str, expected: Any, size: bool = False) -> bool:
    if isinstance(expected, list):
        return any(field_matches(actual, item, size=size) for item in expected)
    return normalize_size(actual) == normalize_size(expected) if size else key(actual) == key(expected)


def rule_matches(rule: dict[str, Any], row: dict[str, str]) -> bool:
    match = rule.get("match", {})
    mapping = {
        "ref_number": "Ref#", "style": "Style", "division": "Division",
        "label": "Label", "product_category": "_product_category",
        "size_configuration": "Size Configuration",
    }
    for rule_field, row_field in mapping.items():
        if rule_field in match and not field_matches(
            row.get(row_field, ""), match[rule_field], size=rule_field == "size_configuration"
        ):
            return False
    if "style_any" in match and not field_matches(row.get("Style", ""), match["style_any"]):
        return False
    if "size_any" in match and not field_matches(row.get("Size Configuration", ""), match["size_any"], size=True):
        return False
    return True


CATEGORY_TERMS = [
    ("BOTTOMS", ("legging", "pant", "short", "skirt", "skort", "jogger", "boardshort")),
    # "crew" is deliberately absent: it names both a crew-neck top and a crew
    # sock, so it misfiles hosiery as a top wherever HOSIERY is not allowed.
    ("TOPS", ("shirt", "top", "tee", "polo", "blouse", "hoody", "hoodie", "sweatshirt")),
    ("SETS", ("set", "2pc", "2 pc", "3pc", "3 pc", "two piece", "multi-piece")),
    ("COVERALLS", ("coverall", "romper")),
    ("BODYSUITS", ("bodysuit",)),
    ("DRESSES", ("dress",)),
    ("OUTERWEAR", ("jacket", "coat", "vest", "outerwear")),
    ("HOSIERY", ("sock", "hosiery", "bib", "blanket")),
    ("UNDERWEAR", ("underwear", "brief", "boxer")),
    ("COLD WEATHER", ("beanie", "glove", "mitten", "scarf", "hat", "cap")),
    ("BAGS", ("bag", "backpack")),
    ("GIRLS SWIMWEAR", ("swim", "bikini", "swimsuit")),
]


def infer_category(description: str, allowed: set[str]) -> tuple[str, str]:
    """Conservative deterministic classification; unresolved text stays unresolved.

    The head of an English compound noun is its last word, so "KNIT SHORT SET"
    is a set rather than a bottom. Scoring by mere presence instead deadlocked on
    every multi-piece description, which is most of a children's apparel order.
    """
    text = key(description)
    matches = []
    for category, words in CATEGORY_TERMS:
        if category not in allowed:
            continue
        head = max((text.rfind(word) for word in words if word in text), default=-1)
        if head >= 0:
            matches.append((head, category))
    if matches:
        last = max(position for position, _ in matches)
        winners = [category for position, category in matches if position == last]
        if len(winners) == 1:
            return winners[0], f"Classified from Product Description: {description}"
    return "", f"Product category is unresolved from Product Description: {description or '[blank]'}"


def parse_source_ranges(value: Any) -> list[str]:
    """Normalize one or several A1 ranges into an ordered, deduplicated list.

    Accepts a list, or a string holding comma-separated ranges. Returns an empty
    list when any part is not a well-formed range, so a malformed citation can
    never be read as a narrower valid one.
    """
    parts = (
        [clean(item).upper() for item in value]
        if isinstance(value, (list, tuple))
        else [part.strip() for part in clean(value).upper().split(",")]
    )
    ranges: list[str] = []
    for part in parts:
        if not part:
            continue
        if not SOURCE_RANGE_RE.fullmatch(part):
            return []
        if part not in ranges:
            ranges.append(part)
    return ranges


def validate_source(sof_wb, source: dict[str, Any]) -> tuple[bool, str]:
    sheet = clean(source.get("sheet"))
    raw_cells = source.get("cells")
    ranges = parse_source_ranges(raw_cells)
    if not sheet or sheet not in sof_wb.sheetnames:
        return False, f"Source sheet is missing or does not exist: {sheet or '[blank]'}"
    if not ranges:
        return False, f"Source cell range is invalid: {clean(raw_cells) or '[blank]'}"
    if len(ranges) > MAX_SOURCE_RANGES:
        return False, (
            f"At most {MAX_SOURCE_RANGES} source ranges may be cited, got {len(ranges)}: "
            + ", ".join(ranges)
        )
    ws = sof_wb[sheet]
    for cells in ranges:
        try:
            min_col, min_row, max_col, max_row = range_boundaries(cells.replace("$", ""))
        except ValueError:
            return False, f"Source cell range is invalid: {cells}"
        if max_row > ws.max_row or max_col > ws.max_column:
            return False, f"Source range {sheet}!{cells} lies outside the used worksheet"
        if not any(clean(ws.cell(r, c).value) for r in range(min_row, max_row + 1)
                   for c in range(min_col, max_col + 1)):
            return False, f"Source range {sheet}!{cells} contains no evidence"
    return True, ""


def validate_rule_result(rule: dict[str, Any], allowed: set[str]) -> tuple[bool, str]:
    result = rule.get("result", {})
    required = [
        "product_category", "sof_hang_flat", "hanger_code", "hanger_color",
        "color_sizer", "sticker_hanger", "size_sticker_hanger",
    ]
    missing = [name for name in required if not clean(result.get(name))]
    if missing:
        return False, "Required SOF data is blank: " + ", ".join(missing)
    not_stated = [name for name in required if key(result.get(name)) == key(NOT_STATED)]
    if not_stated:
        return False, "SOF does not state: " + ", ".join(not_stated)
    if clean(result["product_category"]).upper() not in allowed:
        return False, f"Product category is not allowed: {result['product_category']}"
    hang_flat = normalize_hang_flat(result["sof_hang_flat"])
    if hang_flat not in VALID_HANG_FLAT:
        return False, f"SOF Hang/Flat is invalid: {result['sof_hang_flat']}"
    if hang_flat == "Flat":
        for name in ("hanger_code", "hanger_color", "color_sizer", "sticker_hanger", "size_sticker_hanger"):
            if key(result[name]) != "no":
                return False, f"Flat rule must set {name} to NO"
    # A hung garment always hangs on a physical hanger, so a blank-equivalent
    # code or colour means the value was not found, not that none is required.
    if hang_flat == "Hang":
        for name in ("hanger_code", "hanger_color"):
            if key(result[name]) == "no":
                return False, f"Hang rule must state a real {name}, not NO"
    return True, ""


def resolve_row(row: dict[str, str], rules: list[dict[str, Any]], allowed: set[str], sof_wb,
                selected_sof: Path) -> dict[str, Any]:
    category, category_note = infer_category(row.get("Product Description", ""), allowed)
    row = dict(row)
    row["_product_category"] = category
    candidates = []
    invalid_notes = []
    for rule in rules:
        priority, method = rule_priority(rule.get("match", {}))
        if priority == 99 or not rule_matches(rule, row):
            continue
        valid, note = validate_rule_result(rule, allowed)
        if valid:
            valid, note = validate_source(sof_wb, rule.get("source", {}))
        if not valid:
            invalid_notes.append(f"Rule {rule.get('id', '[unnamed]')}: {note}")
            continue
        candidates.append((priority, method, rule))
    base = {
        "product_category": category,
        "sof_hang_flat": "",
        "hanger_code": "", "hanger_color": "", "color_sizer": "",
        "sticker_hanger": "", "size_sticker_hanger": "",
        "source_sheet": "", "source_cells": "", "source_section": "",
        "source_lines": "", "source_excerpt": "", "match_method": "",
        "confidence": 0.0, "status": "REVIEW", "validation_note": "",
    }
    if not candidates:
        base["validation_note"] = "; ".join(invalid_notes) or category_note or "No matching SOF rule"
        if not invalid_notes and category:
            base["validation_note"] = "No matching SOF rule after applying rule priorities"
        return base
    best_priority = min(item[0] for item in candidates)
    best = [item for item in candidates if item[0] == best_priority]
    fingerprints = {
        json.dumps(item[2].get("result", {}), sort_keys=True, ensure_ascii=False)
        + json.dumps(item[2].get("source", {}), sort_keys=True, ensure_ascii=False)
        for item in best
    }
    if len(fingerprints) > 1:
        ids = ", ".join(clean(item[2].get("id")) or "[unnamed]" for item in best)
        base["validation_note"] = f"Multiple SOF rules remain valid at the same priority: {ids}"
        return base
    _, method, rule = best[0]
    result = rule["result"]
    source = rule["source"]
    sof_hf = normalize_hang_flat(result["sof_hang_flat"]) or ""
    order_hf = normalize_hang_flat(row.get("Hang/Flat"))
    base.update({
        "product_category": clean(result["product_category"]).upper(),
        "sof_hang_flat": sof_hf,
        "hanger_code": clean(result["hanger_code"]),
        "hanger_color": clean(result["hanger_color"]),
        "color_sizer": clean(result["color_sizer"]),
        "sticker_hanger": clean(result["sticker_hanger"]),
        "size_sticker_hanger": clean(result["size_sticker_hanger"]),
        "source_sheet": clean(source["sheet"]), "source_cells": clean(source["cells"]).upper(),
        "match_method": method, "confidence": float(rule.get("confidence", 1.0)),
    })
    citation = f"{selected_sof.name} | {base['source_sheet']}!{base['source_cells']}"
    if order_hf is None:
        base["status"] = "REVIEW"
        base["validation_note"] = f"Missing or unrecognized order Hang/Flat value. Source: {citation}."
    elif order_hf != sof_hf:
        base["status"] = "MISMATCH"
        base["validation_note"] = (
            f"MISMATCH: Order states {order_hf}, but SOF {row.get('Account')} requires {sof_hf}. "
            f"Source: {base['source_sheet']}!{base['source_cells']}."
        )
    else:
        base["status"] = "MATCHED"
    return base


def unresolved_text_row(row: dict[str, str], allowed: set[str]) -> dict[str, Any]:
    category, note = infer_category(row.get("Product Description", ""), allowed)
    return {
        "product_category": category, "sof_hang_flat": "",
        "hanger_code": "", "hanger_color": "", "color_sizer": "",
        "sticker_hanger": "", "size_sticker_hanger": "",
        "source_sheet": "", "source_cells": "", "source_section": "",
        "source_lines": "", "source_excerpt": "", "match_method": "",
        "confidence": 0.0, "status": "REVIEW",
        "validation_note": f"Text SOF requires DeepSeek verification; {note}",
    }


def unmatched_document_row(row: dict[str, str], allowed: set[str], order_format: str,
                           missing: list[str]) -> dict[str, Any]:
    """Report a Word/PDF order row that carries no account.

    An Excel order row without an account belongs to a run whose SOF simply was
    not uploaded, so it is out of scope. A Word/PDF purchase order is the whole
    upload, so the same row is not out of scope: the document did not state the
    account. It is reviewed, with the columns the document lacks named, rather
    than skipped silently or guessed.
    """
    category, note = infer_category(row.get("Product Description", ""), allowed)
    lacking = ", ".join(missing) if missing else "Account"
    return {
        "product_category": category, "sof_hang_flat": "",
        "hanger_code": "", "hanger_color": "", "color_sizer": "",
        "sticker_hanger": "", "size_sticker_hanger": "",
        "source_sheet": "", "source_cells": "", "source_section": "",
        "source_lines": "", "source_excerpt": "", "match_method": "",
        "confidence": 0.0, "status": "REVIEW",
        "validation_note": (
            f"The {order_format.upper()} order states no Account, so no SOF could be selected. "
            f"Columns the document does not carry: {lacking}. {note}"
        ),
    }


def cited_source_values(sof_wb, sheet: str, cells: Any) -> list[str]:
    """Return each non-empty cell of a validated citation as its own statement."""
    ws = sof_wb[sheet]
    values: list[str] = []
    for rng in parse_source_ranges(cells):
        min_col, min_row, max_col, max_row = range_boundaries(rng.replace("$", ""))
        values.extend(
            clean(ws.cell(row, col).value)
            for row in range(min_row, max_row + 1)
            for col in range(min_col, max_col + 1)
            if clean(ws.cell(row, col).value)
        )
    return values


def cited_source_text(sof_wb, sheet: str, cells: Any) -> str:
    """Return normalized text from a previously validated SOF citation."""
    return " ".join(cited_source_values(sof_wb, sheet, cells))


# Words that name each output field in SOF prose. A "NO" claim has to be denied
# next to its own field's wording: scanning the whole cited range instead let an
# unrelated "N/A" in a neighbouring table license a blank value for any field.
# "size sticker" is kept apart from "sticker": a brand-sticker instruction says
# nothing about which size sticker to use, and treating them as one term let any
# brand-sticker paragraph veto a legitimate NO.
FIELD_EVIDENCE_TERMS = {
    "hanger_code": ("hanger",),
    "hanger_color": ("hanger",),
    "color_sizer": ("sizer", "size clip", "crown"),
    "sticker_hanger": ("sticker",),
    "size_sticker_hanger": ("size sticker",),
}
NEGATION_RE = re.compile(r"(?<![a-z0-9])(?:no|n/a|none|not required|without)(?![a-z0-9])")


def no_claim_is_supported(field: str, statements: list[str]) -> bool:
    """Decide whether "NO" is evidence of absence for one field.

    A SOF names an accessory only where it is required, so silence about one is
    absence. Once a cell does discuss it, only a denial inside that same cell
    means NO. The test is per cell because the cited range is flattened for the
    other checks, and across a flattened join an unrelated "N/A" from one table
    row lands next to another row's wording.
    """
    terms = FIELD_EVIDENCE_TERMS.get(field, ())
    mentions = [key(item) for item in statements if any(term in key(item) for term in terms)]
    if not mentions:
        return True
    return any(NEGATION_RE.search(text) for text in mentions)


HANGER_COLOURS = ("black", "white")


def colour_stated_with_code(statements: list[str], hanger_code: str) -> str:
    """The hanger colour named in the same cell as the code, if the SOF names one.

    A cell like "485B/6010B Black Crown Sizer w/ White Lettering" describes the
    hanger first and the lettering second, so the leading colour is the hanger's.
    A cell holding the bare code names no colour and leaves the field to a
    general instruction elsewhere on the sheet.
    """
    code = key(hanger_code)
    if not code:
        return ""
    for text in statements:
        lowered = key(text)
        if code not in lowered:
            continue
        named = [colour for colour in HANGER_COLOURS if colour in lowered]
        if named:
            return min(named, key=lowered.index)
    return ""


def claim_is_supported(claim: str, statements: list[str], field: str) -> bool:
    normalized_claim = key(claim)
    normalized_evidence = key(" ".join(statements))
    if normalized_claim == "no":
        if normalized_evidence == "no":
            return True
        return no_claim_is_supported(field, statements)
    if not normalized_claim:
        return False
    if normalized_claim in normalized_evidence:
        return True
    tokens = [
        token for token in re.findall(r"[a-z0-9]+", normalized_claim)
        if len(token) >= 3 and token not in {"and", "the", "with", "for", "size", "hanger"}
    ]
    if not tokens:
        return False
    matches = sum(token in normalized_evidence for token in set(tokens))
    return matches >= min(2, len(set(tokens)))


def identifier_is_cited(identifier: str, evidence: str) -> bool:
    """Match identifiers such as G1 or AK as complete tokens, not substrings."""
    normalized = key(identifier)
    if not normalized:
        return False
    return re.search(
        rf"(?<![a-z0-9]){re.escape(normalized)}(?![a-z0-9])", key(evidence)
    ) is not None


def validate_llm_evidence(
    sof_wb, decision: dict[str, Any], row: dict[str, str]
) -> tuple[bool, str]:
    source = {
        "sheet": clean(decision.get("source_sheet")),
        "cells": decision.get("source_cells"),
    }
    valid, note = validate_source(sof_wb, source)
    if not valid:
        return False, note
    source_sheet = clean(source["sheet"]).upper()
    if source_sheet not in LLM_SOURCE_SHEETS:
        return False, (
            f"Cited sheet '{source['sheet']}' holds general packing rules, not a "
            f"per-category hanger specification"
        )
    statements = cited_source_values(sof_wb, source["sheet"], source["cells"])
    evidence = " ".join(statements)
    if source_sheet == "PACKAGING":
        label = clean(row.get("Label"))
        if not label:
            return False, "PACKAGING evidence requires a non-blank order Label"
        if not identifier_is_cited(label, evidence):
            return False, (
                f"Cited PACKAGING range does not contain order Label {label}"
            )
        evidence_key = key(evidence)
        if "flatpack" not in evidence_key and "hanger" not in evidence_key:
            return False, "Cited PACKAGING range has no Flat/Hanger requirement"
    sof_hf = normalize_hang_flat(decision.get("sof_hang_flat"))
    evidence_key = key(evidence)
    if sof_hf == "Hang" and "hang" not in evidence_key:
        return False, "Cited SOF range does not support a Hang requirement"
    if sof_hf == "Flat" and not any(
        marker in evidence_key for marker in ("flat", "no hanger", "hanger required no", "not required")
    ):
        return False, "Cited SOF range does not support a Flat requirement"
    # A Flat requirement denies every hanger field at once, and the evidence for
    # it was just checked, so the per-field denials need no separate citation.
    if sof_hf == "Flat":
        return True, ""
    # A general "use white hangers" sentence must not override the size column
    # the order falls in, so the code's own cell wins wherever it names a colour.
    code = clean(decision.get("hanger_code"))
    stated_colour = colour_stated_with_code(statements, code)
    claimed_colour = key(decision.get("hanger_color"))
    if stated_colour and claimed_colour != stated_colour:
        return False, (
            f"Cited cell for hanger_code {code} states a {stated_colour.upper()} hanger, "
            f"not {clean(decision.get('hanger_color')) or '[blank]'}"
        )
    for field in (
        "hanger_code", "hanger_color", "color_sizer", "sticker_hanger", "size_sticker_hanger",
    ):
        claim = clean(decision.get(field))
        if not claim_is_supported(claim, statements, field):
            return False, f"Cited SOF range does not support {field}: {claim or '[blank]'}"
    return True, ""


def validate_text_llm_evidence(
    sof_text, decision: dict[str, Any], row: dict[str, str]
) -> tuple[bool, str]:
    section = clean(decision.get("source_section"))
    lines = clean(decision.get("source_lines")).upper()
    quote = clean(decision.get("source_quote"))
    try:
        evidence = sof_text.cited_text(section, lines)
    except AutomationError as exc:
        return False, str(exc)
    if not quote or len(quote) > 500 or key(quote) not in key(evidence):
        return False, "Cited quote is missing or does not occur in the cited SOF lines"
    label = clean(row.get("Label"))
    if label and not identifier_is_cited(label, evidence):
        return False, f"Cited SOF lines do not contain order Label {label}"
    evidence_key = key(evidence)
    category_terms = {
        "BOTTOMS": ("bottoms", "legging", "pant", "short", "skirt", "jogger"),
        "TOPS": ("tops", "shirt", "tee", "polo", "blouse"),
        "SETS": ("sets", "set", "2pc", "3pc"),
        "DRESSES": ("dresses", "dress"),
        "OUTERWEAR": ("outerwear", "jacket", "coat", "vest"),
        "HOSIERY": ("hosiery", "sock", "bib"),
    }
    category = clean(decision.get("product_category")).upper()
    terms = category_terms.get(category, (category.casefold(),))
    if not any(re.search(rf"\b{re.escape(term)}\w*\b", evidence_key) for term in terms):
        return False, f"Cited SOF lines do not support product category {category}"
    sof_hf = normalize_hang_flat(decision.get("sof_hang_flat"))
    # HANGTAG is not garment HANGER evidence.
    if sof_hf == "Hang" and not re.search(r"\b(?:hanger|hangers|hanging|hang)\b", evidence_key):
        return False, "Cited SOF lines do not support a garment Hang requirement"
    if sof_hf == "Flat" and not any(
        marker in evidence_key for marker in ("flat", "no hanger", "hanger required no")
    ):
        return False, "Cited SOF lines do not support a Flat requirement"
    if sof_hf == "Flat":
        return True, ""
    for field in (
        "hanger_code", "hanger_color", "color_sizer", "sticker_hanger", "size_sticker_hanger",
    ):
        claim = clean(decision.get(field))
        if field == "hanger_code" and key(claim) != "no" and not identifier_is_cited(claim, evidence):
            return False, f"Cited SOF lines do not contain exact hanger_code {claim}"
        if not claim_is_supported(claim, evidence.splitlines(), field):
            return False, f"Cited SOF lines do not support {field}: {claim or '[blank]'}"
    return True, ""


def resolve_llm_decision(
    row: dict[str, str],
    current: dict[str, Any],
    decision: dict[str, Any] | None,
    allowed: set[str],
    sof_wb,
    selected_sof: Path,
    min_confidence: float,
) -> dict[str, Any]:
    """Validate one untrusted LLM decision and apply the normal Hang/Flat gate."""
    base = dict(current)
    original_note = clean(base.get("validation_note"))
    if not isinstance(decision, dict):
        base["validation_note"] = f"{original_note}; DeepSeek returned no decision".strip("; ")
        return base
    decision = dict(decision)
    for field in ACCESSORY_FIELDS:
        if key(decision.get(field)) == key(NOT_STATED):
            decision[field] = "NO"
    reasoning = clean(decision.get("reasoning"))
    if clean(decision.get("status")).upper() != "MATCHED":
        llm_note = f"DeepSeek requested review: {reasoning or 'insufficient SOF evidence'}"
        base["validation_note"] = "; ".join(filter(None, (original_note, llm_note)))
        return base
    try:
        confidence = float(decision.get("confidence"))
    except (TypeError, ValueError):
        confidence = -1.0
    if not 0.0 <= confidence <= 1.0 or confidence < min_confidence:
        llm_note = (
            f"DeepSeek confidence {confidence if confidence >= 0 else '[invalid]'} "
            f"is below required {min_confidence}"
        )
        base["validation_note"] = "; ".join(filter(None, (original_note, llm_note)))
        return base
    from text_sof import TextSof

    text_sof = isinstance(sof_wb, TextSof)
    candidate_rule = {
        "result": {
            "product_category": clean(decision.get("product_category")).upper(),
            "sof_hang_flat": clean(decision.get("sof_hang_flat")),
            "hanger_code": clean(decision.get("hanger_code")),
            "hanger_color": clean(decision.get("hanger_color")),
            "color_sizer": clean(decision.get("color_sizer")),
            "sticker_hanger": clean(decision.get("sticker_hanger")),
            "size_sticker_hanger": clean(decision.get("size_sticker_hanger")),
        },
        "source": {
            "sheet": clean(decision.get("source_sheet")),
            "cells": ",".join(parse_source_ranges(decision.get("source_cells"))),
        },
    }
    valid, note = validate_rule_result(candidate_rule, allowed)
    if valid:
        valid, note = (
            validate_text_llm_evidence(sof_wb, decision, row) if text_sof
            else validate_llm_evidence(sof_wb, decision, row)
        )
    if not valid:
        base["validation_note"] = "; ".join(filter(None, (
            original_note, f"Rejected DeepSeek decision: {note}"
        )))
        return base

    result = candidate_rule["result"]
    source = candidate_rule["source"]
    sof_hf = normalize_hang_flat(result["sof_hang_flat"]) or ""
    order_hf = normalize_hang_flat(row.get("Hang/Flat"))
    base.update({
        "product_category": result["product_category"],
        "sof_hang_flat": sof_hf,
        "hanger_code": result["hanger_code"],
        "hanger_color": result["hanger_color"],
        "color_sizer": result["color_sizer"],
        "sticker_hanger": result["sticker_hanger"],
        "size_sticker_hanger": result["size_sticker_hanger"],
        "source_sheet": "" if text_sof else source["sheet"],
        "source_cells": "" if text_sof else source["cells"],
        "source_section": clean(decision.get("source_section")) if text_sof else "",
        "source_lines": clean(decision.get("source_lines")).upper() if text_sof else "",
        "source_excerpt": clean(decision.get("source_quote")) if text_sof else "",
        "match_method": "LLM_DEEPSEEK_VERIFIED_SOURCE",
        "confidence": confidence,
        "validation_note": "",
    })
    citation = (
        f"{selected_sof.name} | {base['source_section']} {base['source_lines']}"
        if text_sof else f"{selected_sof.name} | {source['sheet']}!{source['cells']}"
    )
    if order_hf is None:
        base["status"] = "REVIEW"
        base["validation_note"] = f"Missing or unrecognized order Hang/Flat value. Source: {citation}."
    elif order_hf != sof_hf:
        base["status"] = "MISMATCH"
        base["validation_note"] = (
            f"MISMATCH: Order states {order_hf}, but SOF {row.get('Account')} requires {sof_hf}. "
            f"Source: {citation}."
        )
    else:
        base["status"] = "MATCHED"
    return base


def ensure_columns(ws, header_row: int, columns: dict[str, int]) -> dict[str, int]:
    result = dict(columns)
    fill = PatternFill("solid", fgColor=SOF_HEADER_FILL)
    for field in WRITE_FIELDS:
        if field not in result:
            col = ws.max_column + 1
            ws.cell(header_row, col, field)
            result[field] = col
        # Mark every SOF-sourced column yellow, including ones the order file
        # already carried, so the source of each column stays obvious.
        ws.cell(header_row, result[field]).fill = fill
    return result


def prepare_review_sheet(wb):
    if "HANGER REVIEW" in wb.sheetnames:
        del wb["HANGER REVIEW"]
    ws = wb.create_sheet("HANGER REVIEW")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(REVIEW_HEADERS))}1"
    fill = PatternFill("solid", fgColor="1F4E78")
    for col, header in enumerate(REVIEW_HEADERS, 1):
        cell = ws.cell(1, col, header)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
    return ws


def append_review(ws, order_sheet: str, row_number: int, row: dict[str, str], result: dict[str, Any], sof_name: str):
    source = ""
    if result.get("source_sheet") and result.get("source_cells"):
        source = f"{sof_name} | {result['source_sheet']}!{result['source_cells']}"
    elif result.get("source_section") and result.get("source_lines"):
        source = f"{sof_name} | {result['source_section']} {result['source_lines']}"
    values = [
        order_sheet, row_number, row.get("PO#"), row.get("Account"), row.get("Division"),
        row.get("Label"), row.get("Label Name"), row.get("Ref#"), row.get("Style"),
        row.get("Product Description"),
        result.get("product_category"), row.get("Size Configuration"), row.get("Hang/Flat"),
        result.get("sof_hang_flat"), result.get("hanger_code"), result.get("hanger_color"),
        result.get("color_sizer"), result.get("sticker_hanger"), result.get("size_sticker_hanger"),
        result.get("status"), result.get("validation_note"), source, result.get("match_method"),
        result.get("confidence"),
    ]
    ws.append(values)


def audit_row_result(row_number: int, row: dict[str, str], result: dict[str, Any],
                     source_file: str) -> dict[str, Any]:
    """Return the stable per-row contract consumed by n8n or other systems."""
    normalized_order_hf = normalize_hang_flat(row.get("Hang/Flat"))
    return {
        "row_number": row_number,
        "po_number": row.get("PO#", ""),
        "account": row.get("Account", ""),
        "division": row.get("Division", ""),
        "label": row.get("Label", ""),
        "label_name": row.get("Label Name", ""),
        "ref_number": row.get("Ref#", ""),
        "style": row.get("Style", ""),
        "product_description": row.get("Product Description", ""),
        "product_category": result.get("product_category", ""),
        "size_configuration": row.get("Size Configuration", ""),
        "order_hang_flat": normalized_order_hf or row.get("Hang/Flat", ""),
        "sof_hang_flat": result.get("sof_hang_flat", ""),
        "hanger_code": result.get("hanger_code", ""),
        "hanger_color": result.get("hanger_color", ""),
        "color_sizer": result.get("color_sizer", ""),
        "sticker_hanger": result.get("sticker_hanger", ""),
        "size_sticker_hanger": result.get("size_sticker_hanger", ""),
        "source_file": source_file,
        "source_sheet": result.get("source_sheet", ""),
        "source_cells": result.get("source_cells", ""),
        "source_section": result.get("source_section", ""),
        "source_lines": result.get("source_lines", ""),
        "source_excerpt": result.get("source_excerpt", ""),
        "match_method": result.get("match_method", ""),
        "confidence": result.get("confidence", 0.0),
        "status": result.get("status", "REVIEW"),
        "validation_note": result.get("validation_note", ""),
        "original_values": {field: row.get(field, "") for field in ORDER_FIELDS},
    }


def build_result_workbook(row_results: list[dict[str, Any]], target: Path) -> Path:
    """Write the one-sheet result deliverable with TONG HOP header colours.

    Hanger values are written only for MATCHED rows; MISMATCH and REVIEW rows keep
    the SOF block empty so an unverified value can never be mistaken for a checked
    one.
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "KET QUA"
    audit_names = [name for name, _ in RESULT_AUDIT_COLUMNS]
    headers = RESULT_ORDER_COLUMNS + RESULT_SOF_COLUMNS + audit_names
    fills = (
        [PatternFill("solid", fgColor=ORDER_HEADER_FILL)] * len(RESULT_ORDER_COLUMNS)
        + [PatternFill("solid", fgColor=SOF_HEADER_FILL)] * len(RESULT_SOF_COLUMNS)
        + [PatternFill("solid", fgColor=AUDIT_HEADER_FILL)] * len(audit_names)
    )
    for index, (name, fill) in enumerate(zip(headers, fills), 1):
        cell = ws.cell(1, index, name)
        cell.font = Font(bold=True)
        cell.fill = fill
    for entry in row_results:
        original = entry.get("original_values") or {}
        row = [original.get(name, "") for name in RESULT_ORDER_COLUMNS]
        matched = entry.get("status") == "MATCHED"
        for name in RESULT_SOF_COLUMNS:
            field = RESULT_SOF_SOURCE.get(name)
            row.append(entry.get(field, "") if (matched and field) else "")
        # The audit block is filled for every row: it is what explains a REVIEW.
        row.extend(entry.get(field, "") for _, field in RESULT_AUDIT_COLUMNS)
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    for index, name in enumerate(headers, 1):
        width = 42 if name == "Validation Note" else max(10, min(len(name) + 4, 30))
        ws.column_dimensions[get_column_letter(index)].width = width
    atomic_save(wb, target)
    return target


def output_path_for(input_path: Path, output_dir: Path, now: datetime) -> Path:
    base = input_path.stem
    stem = f"{base}_checked_{now.strftime('%Y%m%d_%H%M%S')}"
    candidate = output_dir / f"{stem}.xlsx"
    if not candidate.exists():
        return candidate
    counter = 2
    while (output_dir / f"{stem}_{counter}.xlsx").exists():
        counter += 1
    return output_dir / f"{stem}_{counter}.xlsx"


def atomic_save(wb, target: Path):
    target.parent.mkdir(parents=True, exist_ok=True)
    temp = target.with_name(f".{target.name}.tmp.xlsx")
    try:
        wb.save(temp)
        os.replace(temp, target)
    except Exception:
        if temp.exists():
            temp.unlink()
        raise


def read_rules(path: Path) -> tuple[dict[str, list[dict[str, Any]]], set[str]]:
    with path.open("r", encoding="utf-8") as handle:
        payload = json.load(handle)
    if payload.get("schema_version") != 1:
        raise AutomationError("Unsupported rule schema version")
    allowed = {clean(item).upper() for item in payload.get("allowed_categories", []) if clean(item)}
    if not allowed:
        raise AutomationError("Rules file has no allowed_categories")
    accounts = {clean(account).upper(): rules for account, rules in payload.get("accounts", {}).items()}
    return accounts, allowed


def run(input_path: Path, sof_root: Path, rules_path: Path, output_dir: Path,
        overwrite: bool = False, order_sheet: str | None = None, header_scan_rows: int = 50,
        now: datetime | None = None, llm_classifier=None) -> dict[str, Any]:
    started = now or datetime.now()
    if not input_path.exists():
        raise AutomationError(f"Input file does not exist: {input_path}")
    order_format = input_path.suffix.casefold().lstrip(".")
    if order_format not in {"xls", "xlsx", "docx", "pdf"}:
        raise AutomationError("Input file must be .xls, .xlsx, .docx or .pdf")
    account_rules, allowed = read_rules(rules_path)
    audit: dict[str, Any] = {
        "input_file": str(input_path),
        "input_last_modified": datetime.fromtimestamp(input_path.stat().st_mtime).astimezone().isoformat(),
        "sof_files": {}, "total_order_rows": 0, "matched": 0, "mismatch": 0,
        "review": 0, "skipped_existing_hanger_data": 0, "output_file": "",
        "result_file": "", "result_file_name": "", "errors": [],
        "row_results": [], "skipped_rows": [], "skipped_no_sof": 0,
        "accounts_checked": [], "accounts_skipped": [],
        "order_source": {"format": order_format},
        "llm": {"enabled": llm_classifier is not None},
    }
    with tempfile.TemporaryDirectory(prefix="hanger_automation_") as temp_name:
        temp_dir = Path(temp_name)
        text_order = None
        if order_format in {"docx", "pdf"}:
            from text_order import load_text_order

            text_order = load_text_order(input_path, temp_dir)
            audit["order_source"] = {
                "format": text_order.format,
                "extracted_rows": text_order.row_count,
                "fields_found": list(text_order.fields),
                "fields_missing": list(text_order.missing_fields),
                "document_fields": dict(text_order.document_fields),
                "notes": list(text_order.notes),
            }
        wb, _ = load_editable_workbook(text_order.path if text_order else input_path, temp_dir)
        ws, header_row, columns = select_order_sheet(wb, order_sheet, header_scan_rows)
        columns = ensure_columns(ws, header_row, columns)
        review_ws = prepare_review_sheet(wb)
        rows: list[tuple[int, dict[str, str]]] = []
        accounts: set[str] = set()
        for row_number in range(header_row + 1, ws.max_row + 1):
            values = row_values(ws, row_number, columns)
            if not any(values.get(field) for field in ("PO#", "Account", "Ref#", "Style", "Product Description")):
                continue
            rows.append((row_number, values))
            if values["Account"]:
                accounts.add(values["Account"].upper())
        audit["total_order_rows"] = len(rows)
        selections = {account: find_sof(sof_root, account) for account in accounts}
        covered = sorted(a for a, sel in selections.items()
                         if sel.status == "MATCHED" and sel.path)
        audit["accounts_checked"] = covered
        audit["accounts_skipped"] = sorted(set(accounts) - set(covered))
        if not covered:
            audit["errors"].append(
                "No uploaded SOF matched any account in the order file; nothing was checked"
            )
        sof_workbooks = {}
        for account, selection in selections.items():
            audit["sof_files"][account] = {
                "status": selection.status, "selected": str(selection.path) if selection.path else "",
                "revision_date": selection.revision_date, "reason": selection.reason,
                "candidates": list(selection.candidates),
            }
            if selection.path:
                if selection.path.suffix.casefold() in {".pdf", ".docx"}:
                    from text_sof import load_text_sof

                    sof_workbooks[account] = load_text_sof(selection.path)
                else:
                    sof_workbooks[account] = load_sof_workbook(selection.path, temp_dir)
        processed: list[dict[str, Any]] = []
        for row_number, values in rows:
            account = values["Account"].upper()
            selection = selections.get(account)
            existing = [values.get(field, "") for field in WRITE_FIELDS]
            if any(existing) and not overwrite:
                audit["skipped_existing_hanger_data"] += 1
                audit["skipped_rows"].append({
                    "row_number": row_number, "po_number": values.get("PO#", ""),
                    "account": values.get("Account", ""),
                    "reason": "Existing hanger data was preserved",
                })
                continue
            if text_order is not None and not values["Account"]:
                processed.append({
                    "row_number": row_number, "values": values, "account": "",
                    "selection": selection,
                    "result": unmatched_document_row(
                        values, allowed, text_order.format, list(text_order.missing_fields)
                    ),
                })
                continue
            if not selection or selection.status != "MATCHED" or not selection.path:
                # Only the accounts whose SOF was actually uploaded are in scope.
                # Everything else is out of scope for this run, not a review item:
                # counting it as REVIEW would bury the rows a human must really look at.
                audit["skipped_no_sof"] += 1
                audit["skipped_rows"].append({
                    "row_number": row_number, "po_number": values.get("PO#", ""),
                    "account": values.get("Account", ""),
                    "reason": selection.reason if selection else "Account is blank",
                })
                continue
            rules = account_rules.get(account, [])
            from text_sof import TextSof

            result = (
                unresolved_text_row(values, allowed)
                if isinstance(sof_workbooks[account], TextSof)
                else resolve_row(values, rules, allowed, sof_workbooks[account], selection.path)
            )
            processed.append({
                "row_number": row_number,
                "values": values,
                "account": account,
                "selection": selection,
                "result": result,
            })

        if llm_classifier is not None:
            pending_accounts = []
            for account in sorted(accounts):
                selection = selections.get(account)
                if not selection or selection.status != "MATCHED" or not selection.path:
                    continue
                candidates = [
                    (item["row_number"], item["values"])
                    for item in processed
                    if item["account"] == account and item["result"]["status"] == "REVIEW"
                ]
                if candidates:
                    pending_accounts.append((account, selection, candidates))

            def classify_account(job):
                account, selection, candidates = job
                try:
                    return account, llm_classifier.classify(
                        candidates, allowed, sof_workbooks[account], selection.path
                    ), None
                except Exception as exc:
                    return account, None, (
                        f"DeepSeek fallback failed for account {account}: "
                        f"{type(exc).__name__}"
                    )

            # Each account waits on its own DeepSeek calls, so running them one
            # after another made the wall time their sum rather than their max.
            # Only the HTTP work overlaps: every account reads its own workbook,
            # and the classifier caps requests in flight across all of them.
            classified = []
            if len(pending_accounts) == 1:
                classified.append(classify_account(pending_accounts[0]))
            elif pending_accounts:
                with ThreadPoolExecutor(max_workers=len(pending_accounts)) as pool:
                    classified = list(pool.map(classify_account, pending_accounts))

            min_confidence = float(getattr(llm_classifier, "min_confidence", 0.85))
            for account, llm_decisions, error in classified:
                if error:
                    audit["errors"].append(error)
                    continue
                selection = selections[account]
                for item in processed:
                    if item["account"] != account or item["result"]["status"] != "REVIEW":
                        continue
                    item["result"] = resolve_llm_decision(
                        item["values"], item["result"],
                        llm_decisions.get(item["row_number"]), allowed,
                        sof_workbooks[account], selection.path, min_confidence,
                    )
            if hasattr(llm_classifier, "audit_info"):
                audit["llm"] = llm_classifier.audit_info()

        for item in processed:
            row_number = item["row_number"]
            values = item["values"]
            selection = item["selection"]
            result = item["result"]
            status = result["status"]
            audit[status.casefold()] += 1
            audit["row_results"].append(audit_row_result(
                row_number, values, result,
                selection.path.name if selection and selection.path else "",
            ))
            if status == "MATCHED":
                target_values = {
                    "Hanger code": result["hanger_code"], "Hanger color": result["hanger_color"],
                    "Color Sizer": result["color_sizer"], "Sticker hanger": result["sticker_hanger"],
                    "SIZE Sticker hanger": result["size_sticker_hanger"],
                }
                for field, value in target_values.items():
                    ws.cell(row_number, columns[field], value)
            else:
                append_review(review_ws, ws.title, row_number, values, result,
                              selection.path.name if selection and selection.path else "")
        review_ws.auto_filter.ref = (
            f"A1:{get_column_letter(len(REVIEW_HEADERS))}{max(1, review_ws.max_row)}"
        )
        for sof_wb in sof_workbooks.values():
            if hasattr(sof_wb, "close"):
                sof_wb.close()
        output_path = output_path_for(input_path, output_dir, started)
        try:
            atomic_save(wb, output_path)
        except Exception as exc:
            audit["errors"].append(f"Output could not be written: {exc}")
            raise
        finally:
            wb.close()
    audit["total_checked_rows"] = len(audit["row_results"])
    audit["output_file"] = str(output_path)
    result_path = output_path.with_name(output_path.stem + "_KETQUA.xlsx")
    try:
        build_result_workbook(audit["row_results"], result_path)
        audit["result_file"] = str(result_path)
        audit["result_file_name"] = result_path.name
    except Exception as exc:
        audit["errors"].append(f"Result sheet could not be written: {exc}")
        audit["result_file"] = ""
        audit["result_file_name"] = ""
    audit_path = output_path.with_suffix(".audit.json")
    with audit_path.open("w", encoding="utf-8") as handle:
        json.dump(audit, handle, ensure_ascii=False, indent=2)
    audit["audit_file"] = str(audit_path)
    return audit


def extract_sof(path: Path, output: Path | None = None) -> dict[str, Any]:
    """Export cell-addressed SOF evidence for an LLM or a human rule author."""
    with tempfile.TemporaryDirectory(prefix="hanger_extract_") as temp_name:
        wb = load_sof_workbook(path, Path(temp_name))
        payload = {"source_file": path.name, "sheets": []}
        for ws in wb.worksheets:
            merged_lookup = {}
            for merged in ws.merged_cells.ranges:
                for row in ws[merged.coord]:
                    for cell in row:
                        merged_lookup[cell.coordinate] = merged.coord
            cells = []
            for row in ws.iter_rows():
                for cell in row:
                    value = clean(cell.value)
                    if value:
                        item = {"cell": cell.coordinate, "value": value}
                        if cell.coordinate in merged_lookup:
                            item["merged_range"] = merged_lookup[cell.coordinate]
                        cells.append(item)
            payload["sheets"].append({"name": ws.title, "cells": cells})
        wb.close()
    if output:
        output.parent.mkdir(parents=True, exist_ok=True)
        with output.open("w", encoding="utf-8") as handle:
            json.dump(payload, handle, ensure_ascii=False, indent=2)
    return payload


def load_config(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        config = json.load(handle)
    base = path.parent
    for field in ("rules_file", "output_directory"):
        raw = Path(config[field])
        if not raw.is_absolute():
            config[field] = str((base / raw).resolve())
    return config


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    run_parser = sub.add_parser("run", help="validate and update a copy of an order workbook")
    run_parser.add_argument("--input", required=True, type=Path)
    run_parser.add_argument("--config", required=True, type=Path)
    run_parser.add_argument("--overwrite", action="store_true")
    extract_parser = sub.add_parser("extract-sof", help="export cell-addressed SOF evidence as JSON")
    extract_parser.add_argument("--input", required=True, type=Path)
    extract_parser.add_argument("--output", type=Path)
    find_parser = sub.add_parser("find-sof", help="resolve one account to one current SOF")
    find_parser.add_argument("--root", required=True, type=Path)
    find_parser.add_argument("--account", required=True)
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        if args.command == "run":
            config = load_config(args.config)
            audit = run(
                input_path=args.input, sof_root=Path(config["sof_root"]),
                rules_path=Path(config["rules_file"]), output_dir=Path(config["output_directory"]),
                overwrite=args.overwrite or bool(config.get("overwrite_existing_hanger_data", False)),
                order_sheet=config.get("order_sheet"),
                header_scan_rows=int(config.get("header_scan_rows", 50)),
            )
            print(json.dumps(audit, ensure_ascii=False))
        elif args.command == "extract-sof":
            print(json.dumps(extract_sof(args.input, args.output), ensure_ascii=False))
        else:
            selection = find_sof(args.root, args.account)
            print(json.dumps({
                "status": selection.status, "path": str(selection.path) if selection.path else "",
                "revision_date": selection.revision_date, "reason": selection.reason,
                "candidates": selection.candidates,
            }, ensure_ascii=False))
        return 0
    except Exception as exc:
        print(json.dumps({"status": "ERROR", "error": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
