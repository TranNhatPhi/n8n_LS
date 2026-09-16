from __future__ import annotations

import os
import tempfile
from datetime import date
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output" / "docx"
OUT_FILE = OUT_DIR / "Hanger_Automation_Workflow_Trinh_Bay.docx"

NAVY = "17365D"
BLUE = "2F75B5"
PALE_BLUE = "DDEBF7"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F2F2"
MID_GRAY = "D9E2F3"
TEXT = "1F2937"
MUTED = "5B6573"
GREEN = "2E7D32"
PALE_GREEN = "E2F0D9"
AMBER = "B45F06"
PALE_AMBER = "FFF2CC"
RED = "C00000"
PALE_RED = "FCE4D6"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=95, bottom=80, end=95) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9E1F2", size="6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_REPEAT_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_header(row) -> None:
    set_REPEAT_table_header(row)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_repeat_table_header(row) -> None:
    set_repeat_header(row)


def set_col_width(cell, width_in: float) -> None:
    cell.width = Inches(width_in)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_in * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_keep_with_next(paragraph, value=True) -> None:
    paragraph.paragraph_format.keep_with_next = value


def add_text(paragraph, text, *, bold=False, color=TEXT, size=9.5, italic=False):
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def add_bullet(doc, text, level=0, color=TEXT, spacing=2):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(spacing)
    p.paragraph_format.line_spacing = 1.05
    add_text(p, text, color=color, size=9.5)
    return p


def add_number(doc, number, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.08
    p.paragraph_format.left_indent = Inches(0.02)
    p.paragraph_format.first_line_indent = Inches(-0.02)
    add_text(p, f"{number}.  ", color=TEXT, size=9.5)
    add_text(p, title + ": ", bold=True, color=NAVY, size=9.5)
    add_text(p, body, size=9.5)
    return p


def add_callout(doc, title, body, *, fill=LIGHT_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    set_col_width(cell, 6.65)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 120, 160, 120, 160)
    set_cell_border(cell, accent, "10")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    add_text(p, title, bold=True, color=accent, size=10)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.08
    add_text(p2, body, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_table(doc, headers, rows, widths, *, header_fill=NAVY, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for idx, (cell, text) in enumerate(zip(header.cells, headers)):
        set_col_width(cell, widths[idx])
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell, 85, 90, 85, 90)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        add_text(p, str(text), bold=True, color=WHITE, size=font_size)
    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        fill = WHITE if r_idx % 2 == 0 else LIGHT_BLUE
        for idx, (cell, text) in enumerate(zip(cells, row)):
            set_col_width(cell, widths[idx])
            set_cell_shading(cell, fill)
            set_cell_margins(cell, 78, 90, 78, 90)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.02
            add_text(p, str(text), size=font_size)
    return table


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = False
    return p


def para(doc, text, *, bold_lead=None, italic=False, color=TEXT, align=None, after=5, size=9.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        add_text(p, bold_lead, bold=True, color=NAVY, size=size)
        add_text(p, text[len(bold_lead):], color=color, size=size, italic=italic)
    else:
        add_text(p, text, color=color, size=size, italic=italic)
    return p


def font_path(bold=False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial Bold.ttf" if bold else "/Library/Fonts/Arial.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    return next(path for path in candidates if os.path.exists(path))


def pil_color(value: str) -> str:
    return value if value.startswith("#") else f"#{value}"


def draw_centered_text(draw, box, text, font, fill, spacing=4):
    x0, y0, x1, y1 = box
    max_width = x1 - x0 - 24
    words = text.split()
    lines = []
    current = ""
    for word in words:
        test = word if not current else current + " " + word
        if draw.textbbox((0, 0), test, font=font)[2] <= max_width:
            current = test
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    line_h = font.size + spacing
    y = y0 + ((y1 - y0) - line_h * len(lines)) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        x = x0 + ((x1 - x0) - (bbox[2] - bbox[0])) / 2
        draw.text((x, y), line, font=font, fill=pil_color(fill))
        y += line_h


def rounded_box(draw, box, fill, outline, title, subtitle=None):
    draw.rounded_rectangle(box, radius=18, fill=pil_color(fill), outline=pil_color(outline), width=3)
    if subtitle:
        x0, y0, x1, y1 = box
        draw_centered_text(draw, (x0 + 8, y0 + 8, x1 - 8, y0 + 74), title, ImageFont.truetype(font_path(True), 25), NAVY)
        draw_centered_text(draw, (x0 + 12, y0 + 70, x1 - 12, y1 - 8), subtitle, ImageFont.truetype(font_path(False), 18), MUTED)
    else:
        draw_centered_text(draw, box, title, ImageFont.truetype(font_path(True), 22), NAVY)


def arrow(draw, start, end, color=BLUE, width=5):
    draw.line([start, end], fill=pil_color(color), width=width)
    ex, ey = end
    sx, sy = start
    import math
    angle = math.atan2(ey - sy, ex - sx)
    length = 16
    wing = 0.55
    p1 = (ex - length * math.cos(angle - wing), ey - length * math.sin(angle - wing))
    p2 = (ex - length * math.cos(angle + wing), ey - length * math.sin(angle + wing))
    draw.polygon([end, p1, p2], fill=pil_color(color))


def make_workflow_diagram(path: Path):
    image = Image.new("RGB", (2000, 970), pil_color(WHITE))
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(font_path(True), 38)
    lane_font = ImageFont.truetype(font_path(True), 23)
    draw.text((60, 38), "Luồng xử lý Hanger Automation", font=title_font, fill=pil_color(NAVY))
    draw.text((60, 90), "Upload trực tiếp • Rule trước • LLM có kiểm chứng • Không sửa file gốc", font=ImageFont.truetype(font_path(False), 22), fill=pil_color(MUTED))

    lanes = [
        (160, 335, "NGƯỜI DÙNG", "F3F6FA"),
        (350, 595, "N8N ĐIỀU PHỐI", "EAF2F8"),
        (610, 875, "HANGER WORKER", "EEF6EA"),
    ]
    for y0, y1, name, color in lanes:
        draw.rounded_rectangle((45, y0, 1955, y1), radius=20, fill=pil_color(color), outline="#D6DEE8", width=2)
        draw.text((66, y0 + 18), name, font=lane_font, fill=pil_color(NAVY))

    rounded_box(draw, (270, 210, 720, 302), WHITE, BLUE, "1. Tải lên", "1 đơn hàng + 1–20 SOF")
    rounded_box(draw, (820, 210, 1270, 302), WHITE, BLUE, "10. Nhận kết quả", "Excel đã kiểm tra + thông báo")
    arrow(draw, (720, 256), (820, 256))

    boxes_n8n = [
        (145, 425, 430, 535, "2. Kiểm tra", "Định dạng, số lượng, trùng tên"),
        (500, 425, 785, 535, "3. Đóng gói", "Nén các SOF thành ZIP"),
        (855, 425, 1140, 535, "4. Ghép dữ liệu", "Đơn hàng + gói SOF"),
        (1210, 425, 1495, 535, "5. Gọi Worker", "HTTP nội bộ, chờ kết quả"),
        (1565, 425, 1850, 535, "9. Tổng hợp", "SUCCESS hoặc REVIEW"),
    ]
    for box in boxes_n8n:
        rounded_box(draw, box[:4], WHITE, BLUE, box[4], box[5])
    for left, right in zip(boxes_n8n[:4], boxes_n8n[1:4]):
        arrow(draw, (left[2], 480), (right[0], 480))
    arrow(draw, (1495, 480), (1565, 480))

    boxes_worker = [
        (120, 690, 430, 820, "6a. Đọc file", "Excel trực tiếp; Word/PDF lấy chữ"),
        (500, 690, 810, 820, "6b. Rule chuẩn", "Ưu tiên quy tắc đã duyệt"),
        (880, 690, 1190, 820, "6c. DeepSeek", "Chỉ xử lý trường hợp chưa rõ"),
        (1260, 690, 1570, 820, "6d. Kiểm chứng", "Nguồn, trích dẫn, độ tin cậy"),
        (1640, 690, 1880, 820, "7. Xuất file", "MATCHED / MISMATCH / REVIEW"),
    ]
    for box in boxes_worker:
        rounded_box(draw, box[:4], WHITE, GREEN, box[4], box[5])
    for left, right in zip(boxes_worker, boxes_worker[1:]):
        arrow(draw, (left[2], 755), (right[0], 755), color=GREEN)
    arrow(draw, (1352, 535), (275, 690), color=GREEN)
    arrow(draw, (1760, 690), (1707, 535), color=GREEN)
    draw.text((80, 905), "Nguyên tắc an toàn: chỉ ghi dữ liệu hanger khi kết quả đủ bằng chứng; mọi trường hợp còn nghi ngờ được đưa vào REVIEW.", font=ImageFont.truetype(font_path(True), 22), fill=pil_color(RED))
    image.save(path, quality=95)


def configure_document(doc: Document):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.62)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)
    sec.header_distance = Inches(0.28)
    sec.footer_distance = Inches(0.28)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color, before, after in [
        ("Title", 28, "000000", 0, 6),
        ("Subtitle", 13, MUTED, 0, 10),
        ("Heading 1", 17, "000000", 8, 5),
        ("Heading 2", 12, NAVY, 6, 3),
        ("Heading 3", 10, BLUE, 4, 2),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = style_name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ["List Bullet", "List Bullet 2", "List Number"]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(9.5)

    header = sec.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    add_text(hp, "HANGER AUTOMATION  |  TÀI LIỆU QUẢN LÝ", bold=True, color=NAVY, size=7.5)

    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_after = Pt(0)
    add_text(fp, "Nội bộ • Cập nhật 16/09/2026 • Trang ", color=MUTED, size=7.5)
    add_page_field(fp)


def add_cover(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("HA")
    run.font.name = "Arial"
    run.font.size = Pt(34)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    title.paragraph_format.line_spacing = 1.0
    add_text(title, "QUY TRÌNH TỰ ĐỘNG\nKIỂM TRA HANGER", bold=True, color="000000", size=28)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    add_text(subtitle, "Tài liệu trình bày với quản lý", color=MUTED, size=13)

    table = doc.add_table(rows=3, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    entries = [
        ("Mục tiêu", "Tự động đọc đơn hàng và SO Form để xác định hanger"),
        ("Nền tảng", "n8n + Python Worker + DeepSeek, chạy bằng Docker trên MacBook"),
        ("Phạm vi", "Excel / Word / PDF có lớp chữ; upload trực tiếp, không phụ thuộc SMB"),
    ]
    for r_idx, (label, value) in enumerate(entries):
        left, right = table.rows[r_idx].cells
        set_col_width(left, 1.35)
        set_col_width(right, 4.95)
        set_cell_shading(left, NAVY)
        set_cell_shading(right, LIGHT_BLUE if r_idx % 2 == 0 else WHITE)
        for c in (left, right):
            set_cell_margins(c, 120, 130, 120, 130)
            set_cell_border(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        add_text(left.paragraphs[0], label, bold=True, color=WHITE, size=9.5)
        add_text(right.paragraphs[0], value, color=TEXT, size=9.5)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(64)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "Cập nhật ngày 16 tháng 9 năm 2026", color=MUTED, size=9)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p2, "Phiên bản trình bày nội bộ", color=MUTED, size=8.5, italic=True)
    doc.add_page_break()


def build_document():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    heading(doc, "1. Tóm tắt điều hành", 1)
    add_callout(
        doc,
        "Kết luận ngắn",
        "Workflow đã vận hành theo mô hình upload trực tiếp trên MacBook. Người dùng tải 1 file đơn hàng và từ 1 đến 20 SO Form; hệ thống kiểm tra, phân loại hanger, tạo file Excel kết quả và tách riêng các trường hợp cần xem lại. File gốc không bị chỉnh sửa.",
        fill=PALE_GREEN,
        accent=GREEN,
    )
    para(doc, "Giải pháp được thiết kế theo nguyên tắc “rule trước, LLM sau”. Các quy tắc đã được doanh nghiệp phê duyệt được áp dụng trước; DeepSeek chỉ hỗ trợ các trường hợp chưa xác định được. Kết quả LLM tiếp tục bị kiểm chứng bằng nguồn, trích dẫn và ngưỡng tin cậy trước khi được phép ghi vào file.")
    add_bullet(doc, "Đầu vào linh hoạt: đơn hàng và SO Form nhận Excel, Word hoặc PDF có lớp chữ.")
    add_bullet(doc, "Không cần kết nối thư mục mạng SMB; người dùng tải file trực tiếp trên form n8n.")
    add_bullet(doc, "Không dùng OCR ở phiên bản hiện tại; PDF scan chỉ chứa ảnh sẽ bị từ chối rõ ràng.")
    add_bullet(doc, "Mọi trường hợp không đủ bằng chứng được giữ trống và đưa vào REVIEW, không tự đoán.")
    add_bullet(doc, "Kết quả gồm workbook đã kiểm tra, sheet HANGER REVIEW và file audit JSON để truy vết.")

    heading(doc, "2. Phạm vi đầu vào và đầu ra", 1)
    add_table(
        doc,
        ["Hạng mục", "Quy định hiện tại", "Ý nghĩa vận hành"],
        [
            ("File đơn hàng", "Đúng 1 file; .xls, .xlsx, .docx hoặc .pdf", "Nguồn các dòng cần kiểm tra hanger"),
            ("SO Form", "Từ 1 đến 20 file; .xls, .xlsx, .docx hoặc .pdf", "Nguồn quy tắc hanger theo account/label"),
            ("Word/PDF", "Chỉ đọc lớp chữ và bảng; không xử lý hình ảnh", "Nhanh và ổn định; PDF scan cần chuyển thành searchable PDF trước"),
            ("Giới hạn upload", "Tối đa 100 MiB cho một lần gửi", "Ngăn tải quá lớn và bảo vệ worker"),
            ("File kết quả", "Excel có hậu tố _KETQUA.xlsx", "Tải trực tiếp về máy sau khi xử lý"),
            ("Dấu vết kiểm tra", "Sheet HANGER REVIEW + audit JSON", "Giải thích nguồn, phương pháp, độ tin cậy và ghi chú"),
        ],
        [1.15, 2.75, 2.75],
        font_size=8.6,
    )
    para(doc, "Điểm cần nhớ: hệ thống không thay đổi file đơn hàng gốc. Với đơn hàng Word/PDF, dữ liệu được chuẩn hóa thành workbook Excel mới; trường nào tài liệu không thể hiện thì để trống và đánh dấu REVIEW.", bold_lead="Điểm cần nhớ: ")

    doc.add_page_break()
    heading(doc, "3. Luồng xử lý tổng thể", 1)
    with tempfile.TemporaryDirectory() as tmp:
        diagram_path = Path(tmp) / "workflow.png"
        make_workflow_diagram(diagram_path)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        p.add_run().add_picture(str(diagram_path), width=Inches(6.85))

    para(doc, "Sơ đồ trên tách rõ ba vai trò: người dùng gửi/nhận file; n8n điều phối luồng; Hanger Worker chịu trách nhiệm đọc tài liệu và ra quyết định nghiệp vụ. DeepSeek không giao tiếp trực tiếp với n8n và không được tự ghi dữ liệu nếu chưa qua bước kiểm chứng.", italic=True, color=MUTED, size=8.8)

    heading(doc, "4. Chi tiết các node trong n8n", 1)
    add_table(
        doc,
        ["Bước", "Node", "Chức năng và kiểm soát"],
        [
            ("1", "Upload Order và SOF", "Form nhận 1 đơn hàng và tối đa 20 SOF. Hiển thị rõ định dạng được hỗ trợ và nguyên tắc không OCR."),
            ("2", "Validate Upload", "Kiểm tra đủ file, đúng phần mở rộng, số lượng SOF, trùng tên và tạo danh sách trường binary."),
            ("3", "Đóng gói các SOF", "Nén toàn bộ SOF hợp lệ thành sof-files.zip để truyền sang worker trong một request."),
            ("4", "Ghép đơn hàng và SOF", "Ghép file đơn hàng với gói SOF theo cùng một lượt xử lý."),
            ("5", "Gửi File tới Hanger Worker", "POST nội bộ tới /run-upload; thời gian chờ tối đa 60 phút cho lô lớn."),
            ("6", "Evaluate Audit", "Đọc số liệu audit, dừng khi worker lỗi và xác định có cần người dùng xem lại hay không."),
        ],
        [0.48, 1.75, 4.42],
        font_size=8.35,
    )

    doc.add_page_break()
    heading(doc, "4. Chi tiết các node trong n8n (tiếp)", 1)
    add_table(
        doc,
        ["Bước", "Node", "Chức năng và kiểm soát"],
        [
            ("7", "Needs Review?", "Rẽ nhánh dựa trên mismatch, review hoặc kết quả do LLM hỗ trợ cần hậu kiểm."),
            ("8A", "Review Summary", "Thông báo COMPLETED_WITH_REVIEW; nêu số dòng khớp, sai khác, cần xem lại, thiếu SOF và số lần gọi LLM."),
            ("8B", "Success Summary", "Thông báo COMPLETED khi không còn trường hợp cần chú ý."),
            ("9", "Tải File Kết Quả", "Gọi endpoint tải file theo tên an toàn do worker trả về."),
            ("10", "Form Ending", "Trả file binary về trình duyệt và hiển thị bản tóm tắt kết quả."),
        ],
        [0.48, 1.75, 4.42],
        font_size=8.5,
    )
    heading(doc, "Điều kiện nhánh REVIEW", 2)
    para(doc, "n8n chuyển sang nhánh REVIEW nếu có ít nhất một trong ba nhóm: MISMATCH, REVIEW hoặc dòng MATCHED do LLM đề xuất. Việc đưa LLM-matched vào diện chú ý là chủ động: đây không phải kết luận sai, mà là bước hậu kiểm cho tới khi rule nghiệp vụ được chuẩn hóa thêm.")

    heading(doc, "5. Xử lý bên trong Hanger Worker", 1)
    add_number(doc, 1, "Kiểm tra an toàn file", "Xác minh tên file, chữ ký nội dung, ZIP, số file, dung lượng và ngăn đường dẫn nguy hiểm trước khi mở.")
    add_number(doc, 2, "Đọc đơn hàng", "Excel được đọc trực tiếp. Word/PDF được trích chữ và bảng, tìm các cột như Style, Description, PO Qty, Account, Hang/Flat rồi tạo workbook chuẩn hóa.")
    add_number(doc, 3, "Ghép đúng SOF", "Account trên đơn hàng được so với tên SOF. Nếu có nhiều phiên bản, ưu tiên revision/ngày mới nhất có thể xác minh; trường hợp mơ hồ chuyển REVIEW.")
    add_number(doc, 4, "Áp dụng rule đã duyệt", "Tra rules/hanger_rules.json trước để có kết quả ổn định, dễ kiểm toán và không tốn API.")
    add_number(doc, 5, "DeepSeek hỗ trợ", "Chỉ gửi phần chữ và bằng chứng liên quan cho các dòng chưa giải quyết được; các dòng giống nhau được nhóm để giảm số lần gọi.")
    add_number(doc, 6, "Kiểm chứng đầu ra", "Đối chiếu category, Hang/Flat, hanger code/color, nguồn chính xác, trích dẫn và độ tin cậy tối thiểu 0,65.")
    add_number(doc, 7, "Ghi kết quả", "Chỉ MATCHED mới được điền hanger. MISMATCH và REVIEW không ghi giá trị hanger, đồng thời được đưa vào sheet rà soát.")

    doc.add_page_break()
    heading(doc, "6. Cơ chế quyết định và kiểm soát LLM", 1)
    add_table(
        doc,
        ["Trạng thái", "Khi nào xảy ra", "Hành động của hệ thống"],
        [
            ("MATCHED", "Hang/Flat của đơn hàng phù hợp SOF và dữ liệu hanger có bằng chứng hợp lệ", "Điền hanger code, color và phụ kiện; ghi nguồn vào audit"),
            ("MISMATCH", "Thông tin Hang/Flat trên đơn hàng mâu thuẫn với SOF", "Không ghi hanger; chuyển HANGER REVIEW để người dùng quyết định"),
            ("REVIEW", "Thiếu Account/SOF, nguồn mơ hồ, confidence thấp, citation sai hoặc tài liệu không đủ dữ liệu", "Giữ trống; ghi lý do cụ thể để xử lý thủ công"),
        ],
        [1.05, 3.12, 2.48],
        font_size=8.65,
    )
    heading(doc, "Các hàng rào bảo vệ", 2)
    add_bullet(doc, "DeepSeek trả JSON có cấu trúc; hệ thống không nhận câu trả lời tự do để ghi thẳng vào Excel.")
    add_bullet(doc, "Nguồn trích dẫn phải tồn tại đúng sheet/cell hoặc đúng trang/dòng/đoạn chữ trong SOF.")
    add_bullet(doc, "Kết quả dưới ngưỡng 0,65 hoặc không có bằng chứng chính xác bị hạ xuống REVIEW.")
    add_bullet(doc, "Giá trị color/code không được suy đoán từ HANGTAG hoặc nội dung không liên quan.")
    add_bullet(doc, "Nếu DeepSeek hết credit hoặc API lỗi, quy trình vẫn tạo file nhưng giữ các trường hợp chưa giải quyết ở REVIEW.")
    add_callout(
        doc,
        "Vai trò thực tế của AI",
        "LLM là lớp hỗ trợ đọc hiểu tài liệu không đồng nhất, không phải nguồn quyết định duy nhất. Quy tắc đã phê duyệt và bước kiểm chứng trong mã nguồn mới là cơ chế kiểm soát cuối cùng.",
        fill=PALE_AMBER,
        accent=AMBER,
    )

    heading(doc, "7. Dữ liệu đầu ra và cách sử dụng", 1)
    add_table(
        doc,
        ["Đầu ra", "Nội dung", "Người dùng nên làm gì"],
        [
            ("Workbook _KETQUA.xlsx", "Bản đơn hàng đã kiểm tra, có cột hanger và trạng thái", "Dùng các dòng MATCHED; không thay thế file gốc"),
            ("Sheet HANGER REVIEW", "Danh sách MISMATCH/REVIEW kèm ghi chú", "Lọc theo lý do, đối chiếu SOF và xác nhận thủ công"),
            ("Audit JSON", "Nguồn, trích dẫn, phương pháp, confidence, cache/API", "Dùng khi cần truy vết hoặc điều chỉnh rule"),
            ("Thông báo n8n", "Tổng số checked/matched/mismatch/review và tình trạng LLM", "Đánh giá nhanh chất lượng lượt chạy trước khi tải file"),
        ],
        [1.55, 2.6, 2.5],
        font_size=8.45,
    )

    doc.add_page_break()
    heading(doc, "8. Kiến trúc Docker và an toàn vận hành", 1)
    add_table(
        doc,
        ["Thành phần", "Vai trò", "Kiểm soát chính"],
        [
            ("hanger-n8n", "Form và điều phối workflow", "Chỉ bind localhost:5678; tắt Execute Command và đọc/ghi file tùy ý"),
            ("hanger-worker", "Đọc tài liệu, rule, LLM, xuất kết quả", "Chỉ mở cổng 8080 trong mạng Docker, không public ra máy ngoài"),
            ("hanger-redis", "Cache quyết định LLM", "Giảm gọi API; giới hạn 512 MB, cơ chế LRU, không lưu bền"),
            ("n8n_data", "Lưu cấu hình n8n", "Named volume; không xóa bằng docker compose down -v"),
            ("data/output", "Lưu file kết quả", "Bind mount rõ ràng để quản lý và sao lưu"),
        ],
        [1.35, 2.15, 3.15],
        font_size=8.55,
    )
    heading(doc, "Dữ liệu gửi tới DeepSeek", 2)
    para(doc, "Chỉ phần chữ/cell evidence cần thiết cho việc phân loại được gửi đi; không gửi file Excel/Word/PDF nguyên bản. API key chỉ nằm trong biến môi trường của worker, không đặt trong workflow hoặc mã nguồn.")
    heading(doc, "Giới hạn có chủ đích", 2)
    add_bullet(doc, "Không OCR: PDF scan/image-only phải được chuyển sang searchable PDF hoặc Excel/Word trước khi upload.")
    add_bullet(doc, "Tối đa 20 SOF mỗi lần để kiểm soát thời gian và bộ nhớ.")
    add_bullet(doc, "Mỗi lần xử lý dùng thư mục tạm và tự xóa sau khi hoàn tất.")

    heading(doc, "9. Kết quả kiểm thử và hiệu chuẩn", 1)
    add_table(
        doc,
        ["Hạng mục", "Kết quả", "Cách hiểu đúng"],
        [
            ("Kiểm thử tự động", "68 bài test đạt", "Bao phủ parser, rule, worker, bảo mật file và các lỗi hồi quy chính"),
            ("Lượt chạy đầy đủ v7", "4.491 dòng; 2.052 MATCHED; 2.383 REVIEW; 56 MISMATCH; 55 API calls; 0 lỗi", "Bằng chứng hệ thống chạy hết lô khi tài khoản API đủ credit"),
            ("Hiệu chuẩn ngưỡng 0,65", "98% đồng thuận trên 403 dòng có dữ liệu lịch sử so sánh được", "Chỉ là tập con đối chiếu; không phải độ chính xác tổng thể"),
            ("Lượt v8 một phần", "2.695 MATCHED; 69 MISMATCH; 1.727 REVIEW", "Bị dừng một phần do DeepSeek HTTP 402; cần nạp credit để hiệu chuẩn đầy đủ"),
        ],
        [1.55, 2.55, 2.55],
        font_size=8.35,
    )
    para(doc, "Lưu ý: số liệu hiệu chuẩn phản ánh dữ liệu mẫu hiện có và có thể thay đổi khi bổ sung account, label hoặc quy tắc mới. Chỉ nên dùng để đánh giá xu hướng, không coi là cam kết SLA.", italic=True, color=MUTED, size=8.7)

    doc.add_page_break()
    heading(doc, "10. Hạn chế, rủi ro và đề xuất", 1)
    add_table(
        doc,
        ["Vấn đề", "Ảnh hưởng", "Đề xuất quản lý"],
        [
            ("PDF scan không có lớp chữ", "Không thể trích xuất bằng cơ chế hiện tại", "Yêu cầu searchable PDF hoặc bổ sung OCR ở giai đoạn sau"),
            ("Rule một số account chưa thống nhất", "Tăng REVIEW hoặc xung đột với dữ liệu lịch sử", "Chỉ định owner nghiệp vụ phê duyệt code/color chính thức"),
            ("DeepSeek phụ thuộc credit", "Các dòng cần AI sẽ giữ REVIEW khi API 402", "Duy trì quota/cảnh báo credit; thiết lập quy trình fallback thủ công"),
            ("LLM chưa được coi là nguồn duy nhất", "Cần hậu kiểm một phần kết quả MATCHED do LLM", "Pilot theo account; chuyển mẫu ổn định thành rule deterministic"),
            ("Tài liệu đầu vào không đồng nhất", "Thiếu Account/Hang-Flat làm giảm khả năng tự động", "Chuẩn hóa biểu mẫu hoặc bắt buộc trường tối thiểu ở đầu vào"),
        ],
        [1.75, 2.35, 2.55],
        font_size=8.4,
    )

    heading(doc, "Khuyến nghị triển khai", 2)
    add_number(doc, 1, "Giai đoạn 1 — Pilot có kiểm soát", "Chọn 2–3 account phổ biến, duy trì người kiểm tra sheet HANGER REVIEW và lưu lại các trường hợp sai.")
    add_number(doc, 2, "Giai đoạn 2 — Chuẩn hóa rule", "Mỗi mẫu lặp lại đã được nghiệp vụ xác nhận sẽ chuyển thành rule deterministic để giảm chi phí API và REVIEW.")
    add_number(doc, 3, "Giai đoạn 3 — Mở rộng", "Bổ sung OCR nếu nhu cầu PDF scan đủ lớn; theo dõi tỷ lệ REVIEW, mismatch, thời gian xử lý và chi phí trên mỗi lô.")

    add_callout(
        doc,
        "Đề nghị quyết định",
        "Phê duyệt pilot với dữ liệu thực tế, chỉ định một đầu mối nghiệp vụ xác nhận hanger rule và cấp ngân sách/quota DeepSeek đủ cho giai đoạn hiệu chuẩn.",
        fill=PALE_GREEN,
        accent=GREEN,
    )

    heading(doc, "11. Kịch bản trình bày ngắn với sếp", 1)
    para(doc, "“Hệ thống cho phép người dùng tải trực tiếp một đơn hàng và nhiều SO Form lên n8n. n8n kiểm tra file rồi chuyển sang worker để đọc Excel, Word hoặc PDF có chữ. Worker ưu tiên quy tắc đã duyệt; chỉ khi chưa xác định mới nhờ DeepSeek phân tích. Kết quả AI không được ghi thẳng mà phải qua bước kiểm chứng nguồn và độ tin cậy. Dòng chắc chắn được điền vào Excel; dòng mâu thuẫn hoặc chưa rõ được đưa sang HANGER REVIEW. Vì vậy hệ thống giúp giảm thao tác thủ công nhưng vẫn giữ người dùng ở vị trí kiểm soát cuối cùng.”", italic=True, color=NAVY, size=9.5)

    heading(doc, "Các chỉ số nên theo dõi khi pilot", 2)
    add_bullet(doc, "Tỷ lệ MATCHED / MISMATCH / REVIEW theo account.")
    add_bullet(doc, "Tỷ lệ kết quả LLM được người dùng xác nhận đúng.")
    add_bullet(doc, "Thời gian xử lý và số API calls trên mỗi lô.")
    add_bullet(doc, "Số rule mới được chuẩn hóa sau mỗi vòng review.")

    heading(doc, "Nguồn kỹ thuật dùng để lập tài liệu", 2)
    para(doc, "Workflow n8n: n8n/hanger-automation.workflow.json • Worker: src/worker_api.py • Core xử lý: src/hanger_automation.py • LLM: src/deepseek_classifier.py • Parser Word/PDF: src/text_order.py và src/text_sof.py • Docker: compose.yaml • Bộ test: tests/.", color=MUTED, size=8.5)

    core = doc.core_properties
    core.title = "Quy trình tự động kiểm tra Hanger"
    core.subject = "Mô tả workflow n8n, Hanger Worker và cơ chế kiểm soát DeepSeek"
    core.author = "Hanger Automation Project"
    core.keywords = "n8n, hanger, DeepSeek, Docker, SO Form, workflow"
    core.comments = "Tài liệu nội bộ phục vụ trình bày quản lý"

    doc.save(OUT_FILE)
    print(OUT_FILE)


if __name__ == "__main__":
    build_document()
