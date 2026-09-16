import json
import sys
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from hanger_automation import AutomationError, run  # noqa: E402
from text_order import load_text_order  # noqa: E402
from worker_api import (  # noqa: E402
    Upload,
    collect_uploaded_workbooks,
    safe_upload_name,
    validate_upload_bytes,
)


RULES = Path(__file__).resolve().parents[1] / "rules" / "hanger_rules.json"

ITEMS = [
    ("336", "KNIT LEGGING", "Navy 2T", "26M316", "$1,807.68"),
    ("600", "KNIT LEGGING", "Olive 4T", "36M316", "$3,228.00"),
]


def make_po_pdf(path: Path, items=ITEMS, heading="PO number: PO07564") -> None:
    """A purchase order laid out the way the printed PO is: stacked column
    titles, a wrapped second description line, and a totals row."""
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((60, 70), heading, fontsize=9)
    page.insert_text((50, 118), "Qty", fontsize=9)
    page.insert_text((120, 124), "Description", fontsize=9)
    page.insert_text((300, 124), "Style Number", fontsize=9)
    page.insert_text((450, 124), "Amount", fontsize=9)
    page.insert_text((50, 130), "Ordered", fontsize=9)
    bottom = 160
    for quantity, description, wrapped, style, amount in items:
        page.insert_text((52, bottom), quantity, fontsize=9)
        page.insert_text((120, bottom), description, fontsize=9)
        page.insert_text((300, bottom), style, fontsize=9)
        page.insert_text((450, bottom), amount, fontsize=9)
        page.insert_text((120, bottom + 12), wrapped, fontsize=9)
        bottom += 30
    page.insert_text((50, bottom + 10), "Total Units: 936", fontsize=9)
    page.insert_text((450, bottom + 10), "$5,035.68", fontsize=9)
    document.save(path)


def make_order_docx(path: Path, headings=None, rows=None, paragraphs=()) -> None:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    headings = headings or ["Account", "Ref#", "Style", "Description", "Hang/Flat"]
    rows = rows or [["H040M", "NKG-LEM316", "26M316", "KNIT LEGGING", "Hang"]]
    table = document.add_table(rows=1 + len(rows), cols=len(headings))
    for cell, value in zip(table.rows[0].cells, headings):
        cell.text = value
    for row, values in zip(table.rows[1:], rows):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.save(path)


def make_sof(root: Path) -> Path:
    path = root / "H040M Stock Replenishment SO Form 8.31.26.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BOTTOMS"
    sheet["B15"] = "Nike leggings hanger requirement"
    sheet["C16"] = "Hang: 6110 WHITE"
    sheet["D17"] = "Flat styles: no hanger"
    sheet["G18"] = "White size clip/black lettering"
    workbook.save(path)
    return path


def test_pdf_line_items_keep_their_wrapped_description(tmp_path):
    source = tmp_path / "PO07564.pdf"
    make_po_pdf(source)
    order = load_text_order(source, tmp_path)

    assert order.format == "pdf"
    assert order.row_count == 2
    sheet = load_workbook(order.path).active
    headers = [cell.value for cell in sheet[1]]
    rows = [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=2, values_only=True)]
    assert [row["Style"] for row in rows] == ["26M316", "36M316"]
    # The second line of a description belongs to the item above it, not to the
    # Style column it sits next to.
    assert rows[0]["Product Description"] == "KNIT LEGGING Navy 2T"
    assert rows[0]["PO Qty"] == 336
    # Stated once for the whole order, so it is applied to every line item.
    assert rows[1]["PO#"] == "PO07564"
    # A price column names no order field and is dropped instead of guessed.
    assert "Amount" not in headers


def test_pdf_totals_row_is_not_read_as_an_order_line(tmp_path):
    source = tmp_path / "PO07564.pdf"
    make_po_pdf(source)
    order = load_text_order(source, tmp_path)
    descriptions = [
        row[0] for row in load_workbook(order.path).active.iter_rows(min_row=2, values_only=True)
    ]
    assert len(descriptions) == 2


def test_columns_the_pdf_never_states_stay_blank_and_are_reported(tmp_path):
    source = tmp_path / "PO07564.pdf"
    make_po_pdf(source)
    order = load_text_order(source, tmp_path)

    assert "Account" in order.missing_fields
    assert "Hang/Flat" in order.missing_fields
    assert "Style" in order.fields
    sheet = load_workbook(order.path).active
    headers = [cell.value for cell in sheet[1]]
    account = headers.index("Account") + 1
    assert all(sheet.cell(row, account).value in (None, "") for row in range(2, sheet.max_row + 1))


def test_pdf_order_without_an_account_is_reviewed_not_skipped(tmp_path):
    sof_dir = tmp_path / "sof"
    sof_dir.mkdir()
    make_sof(sof_dir)
    source = tmp_path / "PO07564.pdf"
    make_po_pdf(source)

    audit = run(source, sof_dir, RULES, tmp_path / "out")

    # A purchase order is the whole upload, so its rows are never out of scope:
    # they are reviewed with the missing columns named.
    assert audit["review"] == 2
    assert audit["matched"] == 0
    assert audit["skipped_no_sof"] == 0
    assert audit["order_source"]["format"] == "pdf"
    assert audit["order_source"]["extracted_rows"] == 2
    assert "Account" in audit["order_source"]["fields_missing"]
    note = audit["row_results"][0]["validation_note"]
    assert "states no Account" in note
    assert "Hang/Flat" in note
    assert Path(audit["result_file"]).is_file()
    assert Path(audit["output_file"]).name.startswith("PO07564_checked_")


def test_docx_order_completes_the_hanger_check_like_an_excel_order(tmp_path):
    sof_dir = tmp_path / "sof"
    sof_dir.mkdir()
    make_sof(sof_dir)
    source = tmp_path / "PO order.docx"
    make_order_docx(source)

    audit = run(source, sof_dir, RULES, tmp_path / "out")

    assert audit["matched"] == 1
    assert audit["review"] == 0
    assert audit["order_source"]["format"] == "docx"
    result = audit["row_results"][0]
    assert result["account"] == "H040M"
    assert result["hanger_code"] == "6110"
    assert result["source_sheet"] == "BOTTOMS"


def test_docx_account_stated_once_applies_to_every_line(tmp_path):
    sof_dir = tmp_path / "sof"
    sof_dir.mkdir()
    make_sof(sof_dir)
    source = tmp_path / "PO order.docx"
    make_order_docx(
        source,
        headings=["Ref#", "Style", "Description", "Hang/Flat"],
        rows=[["NKG-LEM316", "26M316", "KNIT LEGGING", "Hang"],
              ["NKG-LEM316", "36M316", "KNIT LEGGING", "Hang"]],
        paragraphs=["Account: H040M", "Ship to: Sydney"],
    )

    audit = run(source, sof_dir, RULES, tmp_path / "out")

    assert audit["matched"] == 2
    assert audit["order_source"]["document_fields"]["Account"] == "H040M"


def test_a_field_stated_twice_with_two_values_is_left_blank(tmp_path):
    source = tmp_path / "PO order.docx"
    make_order_docx(
        source,
        headings=["Ref#", "Style", "Description", "Hang/Flat"],
        rows=[["NKG-LEM316", "26M316", "KNIT LEGGING", "Hang"]],
        paragraphs=["Account: H040M", "Account: H494M"],
    )
    order = load_text_order(source, tmp_path)

    assert "Account" not in order.document_fields
    assert "Account" in order.missing_fields
    assert any("stated more than once" in note for note in order.notes)


def test_scanned_order_pdf_is_rejected_without_ocr(tmp_path):
    source = tmp_path / "PO scan.pdf"
    document = pymupdf.open()
    page = document.new_page()
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 100, 100), False)
    pixmap.clear_with(255)
    page.insert_image(pymupdf.Rect(72, 72, 172, 172), pixmap=pixmap)
    document.save(source)

    with pytest.raises(AutomationError, match="OCR is disabled"):
        load_text_order(source, tmp_path)


def test_a_document_without_a_line_item_table_is_refused(tmp_path):
    source = tmp_path / "letter.docx"
    document = Document()
    document.add_paragraph("Please confirm the shipment dates for PO07564.")
    document.save(source)

    with pytest.raises(AutomationError, match="no readable line-item table"):
        load_text_order(source, tmp_path)


def test_excel_orders_keep_the_workbook_path(tmp_path):
    sof_dir = tmp_path / "sof"
    sof_dir.mkdir()
    make_sof(sof_dir)
    source = tmp_path / "orders.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Account", "Ref#", "Style", "Product Description", "Hang/Flat",
                  "Label", "Division", "Size Configuration", "PO#"])
    sheet.append(["H040M", "NKG-LEM316", "26M316", "KNIT LEGGING", "Hang",
                  "WH", "G", "2T - 3T - 4T", "751419"])
    workbook.save(source)

    audit = run(source, sof_dir, RULES, tmp_path / "out")

    assert audit["matched"] == 1
    assert audit["order_source"] == {"format": "xlsx"}


def test_worker_accepts_word_and_pdf_orders(tmp_path):
    source = tmp_path / "PO07564.pdf"
    make_po_pdf(source)
    sof = tmp_path / "H040M SO Form 8.31.26.xlsx"
    Workbook().save(sof)

    assert safe_upload_name("PO07564.pdf", "order_workbook") == "PO07564.pdf"
    validate_upload_bytes(Upload("PO07564.pdf", source.read_bytes()), "order_workbook")
    order, sofs = collect_uploaded_workbooks(
        {"order_workbook": [Upload("PO07564.pdf", source.read_bytes())],
         "sof_workbook": [Upload(sof.name, sof.read_bytes())]},
        20, 100 * 1024 * 1024,
    )
    assert order.filename == "PO07564.pdf"
    assert len(sofs) == 1

    with pytest.raises(AutomationError, match="must be"):
        safe_upload_name("orders.txt", "order_workbook")
    with pytest.raises(AutomationError, match="order_workbook is not a valid .pdf"):
        validate_upload_bytes(Upload("PO07564.pdf", b"not-a-pdf"), "order_workbook")


def test_n8n_form_accepts_word_and_pdf_orders():
    workflow = json.loads(
        (Path(__file__).resolve().parents[1] / "n8n" / "hanger-automation.workflow.json")
        .read_text(encoding="utf-8")
    )
    nodes = {node["name"]: node for node in workflow["nodes"]}
    fields = nodes["Upload Order và SOF"]["parameters"]["formFields"]["values"]
    order_field = next(field for field in fields if field["fieldName"] == "order_workbook")

    assert order_field["acceptFileTypes"] == ".xls,.xlsx,.docx,.pdf"
    assert order_field["multipleFiles"] is False
    validation = nodes["Validate Upload"]["parameters"]["jsCode"]
    assert "(?:xls|xlsx|docx|pdf)$/i.test(String(order.fileName" in validation
    # SOF files keep their own check, on the SOF name rather than the order name.
    assert "(?:xls|xlsx|docx|pdf)$/i.test(name)" in validation


def test_n8n_form_hands_the_result_workbook_back_as_a_download():
    workflow = json.loads(
        (Path(__file__).resolve().parents[1] / "n8n" / "hanger-automation.workflow.json")
        .read_text(encoding="utf-8")
    )
    nodes = {node["name"]: node for node in workflow["nodes"]}
    ending = nodes["Tải File Về Máy"]

    assert ending["type"] == "n8n-nodes-base.form"
    assert ending["parameters"]["operation"] == "completion"
    assert ending["parameters"]["respondWith"] == "returnBinary"
    assert ending["parameters"]["inputDataFieldName"] == "data"
    # The run summary would otherwise be lost, so the ending page carries it.
    assert "result_file_name" in ending["parameters"]["completionMessage"]
    downstream = workflow["connections"]["Tải File Kết Quả"]["main"][0]
    assert [entry["node"] for entry in downstream] == ["Tải File Về Máy"]


def test_form_trigger_keeps_the_response_mode_its_ending_node_requires():
    """The Form Trigger rejects a Respond to Webhook node and errors on
    responseNode without one, so neither may appear in the workflow."""
    workflow = json.loads(
        (Path(__file__).resolve().parents[1] / "n8n" / "hanger-automation.workflow.json")
        .read_text(encoding="utf-8")
    )
    nodes = {node["name"]: node for node in workflow["nodes"]}

    assert nodes["Upload Order và SOF"]["parameters"]["responseMode"] == "lastNode"
    assert not [n for n in workflow["nodes"] if n["type"] == "n8n-nodes-base.respondToWebhook"]
