import sys
import json
import urllib.request
from pathlib import Path

import pymupdf
import pytest
from docx import Document
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from hanger_automation import AutomationError, run, validate_text_llm_evidence  # noqa: E402
from text_sof import load_text_sof  # noqa: E402
from worker_api import Upload, collect_uploaded_workbooks, validate_sof_bytes  # noqa: E402
from deepseek_classifier import DeepSeekClassifier, DeepSeekSettings  # noqa: E402


RULES = Path(__file__).resolve().parents[1] / "rules" / "hanger_rules.json"


def make_order(path: Path, hang_flat="Flat") -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Orders"
    sheet.append(["Account", "Label", "Product Description", "Style", "PO#",
                  "Size Configuration", "Hang/Flat"])
    sheet.append(["H040M", "WH", "KNIT LEGGING", "S1", "P1", "4-6", hang_flat])
    workbook.save(path)


def make_docx(path: Path, text="WH | BOTTOMS LEGGINGS | FLATPACKED | no hanger") -> None:
    document = Document()
    table = document.add_table(rows=1, cols=4)
    for cell, value in zip(table.rows[0].cells, text.split(" | ")):
        cell.text = value
    document.save(path)


def make_pdf(path: Path, text="WH BOTTOMS LEGGINGS FLATPACKED no hanger") -> None:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    document.save(path)


class FlatClassifier:
    min_confidence = 0.85

    def classify(self, rows, allowed, sof_text, selected_sof):
        section = next(iter(sof_text.sections))
        return {number: {
            "status": "MATCHED", "confidence": 0.95,
            "product_category": "BOTTOMS", "sof_hang_flat": "Flat",
            "hanger_code": "NO", "hanger_color": "NO", "color_sizer": "NO",
            "sticker_hanger": "NO", "size_sticker_hanger": "NO",
            "source_section": section, "source_lines": "L0001",
            "source_quote": "FLATPACKED",
        } for number, _ in rows}

    def audit_info(self):
        return {"enabled": True, "provider": "deepseek", "api_calls": 0}


def test_docx_table_text_is_extracted_without_images(tmp_path):
    path = tmp_path / "H040M SOForm.docx"
    make_docx(path)
    source = load_text_sof(path)
    assert source.format == "docx"
    assert source.cited_text("DOCUMENT", "L0001") == (
        "WH | BOTTOMS LEGGINGS | FLATPACKED | no hanger"
    )


def test_pdf_text_layer_is_extracted_by_page(tmp_path):
    path = tmp_path / "H040M SOForm.pdf"
    make_pdf(path)
    source = load_text_sof(path)
    assert source.format == "pdf"
    assert "FLATPACKED" in source.cited_text("PAGE 1", "L0001")


def test_image_only_pdf_is_rejected_without_ocr(tmp_path):
    path = tmp_path / "H040M scanned.pdf"
    document = pymupdf.open()
    page = document.new_page()
    pixmap = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 100, 100), False)
    pixmap.clear_with(255)
    page.insert_image(pymupdf.Rect(72, 72, 172, 172), pixmap=pixmap)
    document.save(path)
    with pytest.raises(AutomationError, match="OCR is disabled"):
        load_text_sof(path)


@pytest.mark.parametrize("suffix", [".docx", ".pdf"])
def test_text_sof_can_be_uploaded_and_verified_without_excel_rule(tmp_path, suffix):
    sof_dir = tmp_path / "sof"
    sof_dir.mkdir()
    source_path = sof_dir / f"H040M SOForm{suffix}"
    (make_docx if suffix == ".docx" else make_pdf)(source_path)
    order = tmp_path / "order.xlsx"
    make_order(order)
    files = {
        "order_workbook": [Upload(order.name, order.read_bytes())],
        "sof_workbook": [Upload(source_path.name, source_path.read_bytes())],
    }
    _, sofs = collect_uploaded_workbooks(files, 20, 100 * 1024 * 1024)
    assert sofs[0].filename == source_path.name
    audit = run(order, sof_dir, RULES, tmp_path / "out", llm_classifier=FlatClassifier())
    assert audit["matched"] == 1
    assert audit["row_results"][0]["source_section"] in {"DOCUMENT", "PAGE 1"}
    assert audit["row_results"][0]["source_lines"] == "L0001"
    assert audit["row_results"][0]["source_sheet"] == ""


def test_text_sof_without_deepseek_stays_review(tmp_path):
    sof_dir = tmp_path / "sof"
    sof_dir.mkdir()
    make_docx(sof_dir / "H040M SOForm.docx")
    order = tmp_path / "order.xlsx"
    make_order(order)
    audit = run(order, sof_dir, RULES, tmp_path / "out")
    assert audit["review"] == 1
    assert audit["matched"] == 0


def test_hangtag_is_not_accepted_as_garment_hanger_evidence():
    from text_sof import TextSof

    source = TextSof({"DOCUMENT": ("WH BOTTOMS LEGGINGS HANGTAG 6110 WHITE",)}, "docx")
    decision = {"source_section": "DOCUMENT", "source_lines": "L0001",
                "source_quote": "HANGTAG", "product_category": "BOTTOMS",
                "sof_hang_flat": "Hang", "hanger_code": "6110", "hanger_color": "WHITE",
                "color_sizer": "NO", "sticker_hanger": "NO", "size_sticker_hanger": "NO"}
    valid, note = validate_text_llm_evidence(source, decision, {"Label": "WH"})
    assert not valid
    assert "garment Hang" in note


def test_bad_pdf_signature_is_rejected():
    with pytest.raises(AutomationError, match="valid .pdf"):
        validate_sof_bytes(Upload("H040M.pdf", b"not-pdf"))


def test_text_llm_request_contains_only_text_and_line_citations(monkeypatch):
    captured = {}

    class Response:
        def __enter__(self):
            return self

        def __exit__(self, *_):
            return False

        def read(self):
            return json.dumps({"choices": [{"message": {"content": '{"decisions": []}'}}]}).encode()

    def fake_urlopen(request, timeout):
        captured.update(json.loads(request.data))
        return Response()

    monkeypatch.setattr(urllib.request, "urlopen", fake_urlopen)
    classifier = DeepSeekClassifier(DeepSeekSettings(api_keys=("test",)))
    classifier._request({}, {"BOTTOMS"}, "H040M.pdf",
                        "[SECTION: PAGE 1]\nL0001 = WH BOTTOMS flat no hanger\n",
                        "test", text_mode=True)
    user = json.loads(captured["messages"][1]["content"])
    assert "sof_text_evidence" in user
    assert "sof_cell_evidence" not in user
    assert "source_lines" in user["output_schema"]["decisions"][0]
    assert "source_cells" not in user["output_schema"]["decisions"][0]
    assert all(isinstance(message["content"], str) for message in captured["messages"])
