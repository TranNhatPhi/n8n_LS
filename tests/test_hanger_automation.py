import io
import json
import sys
import threading
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from hanger_automation import (  # noqa: E402
    find_sof,
    infer_category,
    resolve_llm_decision,
    run,
)
from deepseek_classifier import (  # noqa: E402
    DecisionCache, DeepSeekClassifier, DeepSeekSettings, _evidence_text,
)
from worker_api import (  # noqa: E402
    AutomationError,
    Upload,
    collect_uploaded_workbooks,
    parse_multipart,
    safe_upload_name,
    validate_workbook_bytes,
)


HEADERS = [
    "Division", "Season", "Year", "Ref#", "Style", "Product Description",
    "Color", "Label", "Label Name", "Account", "PO#", "PO Qty",
    "Size Configuration", "Pack Ratio", "Master Box Quantity", "Hang/Flat",
    "Hanger code", "Hanger color", "Color Sizer", "Sticker hanger",
    "SIZE Sticker hanger",
]

LABEL_NAMES = {"WH": "WITH HANGER", "": ""}

CASES = [
    (751419, "26M316", "WH", "2T - 3T - 4T", "Hang", "6110"),
    (751422, "26M316", "WH", "2T - 3T - 4T", "Hang", "6110"),
    (751423, "3XR316", "", "6X", "Flat", "NO"),
    (751424, "34R316", "", "4", "Flat", "NO"),
    (751425, "35R316", "", "5", "Flat", "NO"),
    (751427, "36M316", "WH", "4 - 5 - 6 - 6X", "Hang", "6110"),
    (766452, "36M316", "WH", "4 - 5 - 6 - 6X", "Hang", "6110"),
    (751430, "36R316", "", "6", "Flat", "NO"),
]


def make_sof(root: Path, name: str = "H040M Stock Replenishment SO Form 8.31.26.xlsx") -> Path:
    path = root / name
    wb = Workbook()
    ws = wb.active
    ws.title = "BOTTOMS"
    ws["B15"] = "Nike leggings hanger requirement"
    ws["C16"] = "Hang: 6110 WHITE"
    ws["D17"] = "Flat styles: no hanger"
    ws["G18"] = "White size clip/black lettering"
    wb.save(path)
    return path


def make_hosiery_sof(root: Path) -> Path:
    path = root / "H040M Stock Replenishment SO Form 8.31.26.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "HOSIERY"
    ws["A3"] = "Bibs, Socks"
    ws["A4"] = "Hanger Required"
    ws["B4"] = "No"
    ws["A5"] = "Flat pack"
    wb.save(path)
    return path


def make_hbe_sof(root: Path) -> Path:
    path = root / "H494M - SOFORM - HBE 7.23.xlsx"
    wb = Workbook()
    general = wb.active
    general.title = "GENERAL"
    general["E1"] = "HBE account"
    general["F1"] = "#H494M"
    packaging = wb.create_sheet("PACKAGING")
    packaging["D23"] = "Labels"
    packaging["E23"] = "FLATPACKED/HANGER"
    packaging["D24"] = "AK-KT-KU-KV-KW"
    packaging["E24"] = "FLATPACKED"
    packaging["D25"] = "G1-G2-G3-G4-G5"
    packaging["E25"] = "FLATPACKED/HANGER - HANGER PACKED FOR SETS"
    packaging["G25"] = "Follow US Manual for Hanger codes"
    wb.save(path)
    return path


class FakeClassifier:
    min_confidence = 0.85

    def __init__(self, decision):
        self.decision = decision
        self.calls = 0

    def classify(self, rows, allowed_categories, sof_wb, selected_sof):
        self.calls += 1
        return {row_number: dict(self.decision) for row_number, _ in rows}

    def audit_info(self):
        return {"enabled": True, "provider": "deepseek", "model": "test", "api_calls": self.calls}


def make_order(path: Path, cases=CASES):
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.freeze_panes = "A2"
    ws.append(HEADERS)
    for po, style, label, sizes, hang_flat, _ in cases:
        values = {
            "Division": "G", "Ref#": "NKG-LEM316", "Style": style,
            "Product Description": "KNIT LEGGING", "Label": label,
            "Label Name": LABEL_NAMES.get(label, ""),
            "Account": "H040M", "PO#": po, "Size Configuration": sizes,
            "Hang/Flat": hang_flat,
        }
        ws.append([values.get(header, "") for header in HEADERS])
        ws.cell(ws.max_row, HEADERS.index("PO#") + 1).number_format = "0000000"
    wb.save(path)


def test_all_eight_acceptance_rows(tmp_path):
    sof_root = tmp_path / "sof"
    sof_root.mkdir()
    make_sof(sof_root)
    order = tmp_path / "3. TONG HOP HANGER-2026.xlsx"
    make_order(order)
    original_bytes = order.read_bytes()
    output = tmp_path / "out"
    rules = Path(__file__).resolve().parents[1] / "rules" / "hanger_rules.json"

    audit = run(order, sof_root, rules, output, now=datetime(2026, 9, 14, 12, 30, 45))

    assert audit["matched"] == 8
    assert audit["mismatch"] == 0
    assert audit["review"] == 0
    assert order.read_bytes() == original_bytes
    assert len(audit["row_results"]) == 8
    assert audit["row_results"][0]["po_number"] == "0751419"
    assert audit["row_results"][0]["source_sheet"] == "BOTTOMS"
    assert audit["row_results"][0]["source_cells"] == "B15:G18"
    assert audit["row_results"][0]["status"] == "MATCHED"
    assert Path(audit["output_file"]).name == "3. TONG HOP HANGER-2026_checked_20260914_123045.xlsx"
    wb = load_workbook(audit["output_file"])
    ws = wb["Orders"]
    columns = {cell.value: cell.column for cell in ws[1]}
    for index, case in enumerate(CASES, 2):
        expected_po = f"{case[0]:07d}"
        assert ws.cell(index, columns["PO#"]).number_format == "0000000"
        assert f"{ws.cell(index, columns['PO#']).value:07d}" == expected_po
        assert ws.cell(index, columns["Hanger code"]).value == case[-1]
        assert ws.cell(index, columns["Hanger color"]).value == ("WHITE" if case[-1] == "6110" else "NO")
        assert ws.cell(index, columns["Color Sizer"]).value == (
            "White size clip/black lettering" if case[-1] == "6110" else "NO"
        )
        assert ws.cell(index, columns["Sticker hanger"]).value == "NO"
        assert ws.cell(index, columns["SIZE Sticker hanger"]).value == "NO"
    assert wb["HANGER REVIEW"].max_row == 1
    wb.close()
    with Path(audit["audit_file"]).open(encoding="utf-8") as handle:
        stored = json.load(handle)
    assert stored["total_order_rows"] == 8


def test_both_mismatch_directions_and_blank_are_reviewed(tmp_path):
    sof_root = tmp_path / "sof"
    sof_root.mkdir()
    make_sof(sof_root)
    cases = [
        (751419, "26M316", "WH", "2T - 3T - 4T", "Flat", ""),
        (751423, "3XR316", "", "6X", "Hang", ""),
        (751427, "36M316", "WH", "4 - 5 - 6 - 6X", "", ""),
    ]
    order = tmp_path / "orders.xlsx"
    make_order(order, cases)
    rules = Path(__file__).resolve().parents[1] / "rules" / "hanger_rules.json"

    audit = run(order, sof_root, rules, tmp_path / "out")

    assert audit["matched"] == 0
    assert audit["mismatch"] == 2
    assert audit["review"] == 1
    wb = load_workbook(audit["output_file"])
    review = wb["HANGER REVIEW"]
    status_column = {c.value: c.column for c in review[1]}["Validation Status"]
    statuses = [review.cell(row, status_column).value for row in range(2, 5)]
    assert statuses == ["MISMATCH", "MISMATCH", "REVIEW"]
    wb.close()


def test_ambiguous_latest_sof_is_not_guessed(tmp_path):
    make_sof(tmp_path, "H040M SO Form 8.31.26.xlsx")
    make_sof(tmp_path, "archive H040M SO Form 8.31.26.xlsx")
    selection = find_sof(tmp_path, "H040M")
    assert selection.status == "REVIEW"
    assert selection.path is None
    assert "share the latest revision date" in selection.reason


def test_undated_candidate_blocks_latest_selection(tmp_path):
    make_sof(tmp_path, "H040M SO Form 8.31.26.xlsx")
    make_sof(tmp_path, "H040M current.xlsx")
    selection = find_sof(tmp_path, "H040M")
    assert selection.status == "REVIEW"
    assert "not verifiable" in selection.reason


def test_existing_hanger_data_is_preserved(tmp_path):
    sof_root = tmp_path / "sof"
    sof_root.mkdir()
    make_sof(sof_root)
    order = tmp_path / "orders.xlsx"
    make_order(order, [CASES[0]])
    wb = load_workbook(order)
    ws = wb["Orders"]
    ws.cell(2, HEADERS.index("Hanger code") + 1, "MANUAL-CODE")
    wb.save(order)
    wb.close()
    rules = Path(__file__).resolve().parents[1] / "rules" / "hanger_rules.json"

    audit = run(order, sof_root, rules, tmp_path / "out")

    assert audit["skipped_existing_hanger_data"] == 1
    assert audit["row_results"] == []
    wb = load_workbook(audit["output_file"])
    assert wb["Orders"].cell(2, HEADERS.index("Hanger code") + 1).value == "MANUAL-CODE"
    wb.close()


def test_uncited_rule_becomes_review(tmp_path):
    sof_root = tmp_path / "sof"
    sof_root.mkdir()
    sof = sof_root / "H040M SO Form 8.31.26.xlsx"
    Workbook().save(sof)
    order = tmp_path / "orders.xlsx"
    make_order(order, [CASES[0]])
    rules = Path(__file__).resolve().parents[1] / "rules" / "hanger_rules.json"

    audit = run(order, sof_root, rules, tmp_path / "out")

    assert audit["review"] == 1
    assert "Source sheet is missing" in audit["row_results"][0]["validation_note"]


def test_worker_sanitizes_and_validates_uploaded_workbooks(tmp_path):
    assert safe_upload_name(r"C:\\fakepath\\orders.xlsx", "order") == "orders.xlsx"
    valid_path = tmp_path / "valid.xlsx"
    Workbook().save(valid_path)
    validate_workbook_bytes(Upload("valid.xlsx", valid_path.read_bytes()), "order")
    # Orders accept Word/PDF as well; anything outside the four formats does not.
    assert safe_upload_name("orders.pdf", "order") == "orders.pdf"
    try:
        safe_upload_name("orders.txt", "order")
    except AutomationError as exc:
        assert "must be .docx, .pdf, .xls, .xlsx" in str(exc)
    else:
        raise AssertionError("Worker accepted an unsupported upload")


def test_worker_parses_n8n_multipart_upload(tmp_path):
    workbook = tmp_path / "orders.xlsx"
    Workbook().save(workbook)
    content = workbook.read_bytes()
    boundary = "hanger-test-boundary"
    parts = []
    for field, filename in (
        ("order_workbook", "orders.xlsx"),
        ("sof_workbook", "H040M SO Form 8.31.26.xlsx"),
    ):
        parts.append(
            f"--{boundary}\r\n"
            f'Content-Disposition: form-data; name="{field}"; filename="{filename}"\r\n'
            "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n"
        .encode() + content + b"\r\n")
    parts.append(
        f"--{boundary}\r\nContent-Disposition: form-data; name=\"overwrite\"\r\n\r\nfalse\r\n"
        f"--{boundary}--\r\n".encode()
    )
    files, fields = parse_multipart(b"".join(parts), f"multipart/form-data; boundary={boundary}")
    assert set(files) == {"order_workbook", "sof_workbook"}
    assert len(files["order_workbook"]) == 1
    assert len(files["sof_workbook"]) == 1
    assert fields["overwrite"] == "false"


def test_worker_accepts_multiple_n8n_sof_fields(tmp_path):
    workbook = tmp_path / "valid.xlsx"
    Workbook().save(workbook)
    content = workbook.read_bytes()
    files = {
        "order_workbook": [Upload("orders.xlsx", content)],
        "sof_workbook_0": [Upload("H040M SO Form.xlsx", content)],
        "sof_workbook_1": [Upload("Q777Q SO Form.xlsx", content)],
    }

    order, sofs = collect_uploaded_workbooks(files, 20, 100 * 1024 * 1024)

    assert order.filename == "orders.xlsx"
    assert [upload.filename for upload in sofs] == [
        "H040M SO Form.xlsx", "Q777Q SO Form.xlsx",
    ]


def test_worker_accepts_n8n_generated_sof_zip(tmp_path):
    workbook = tmp_path / "valid.xlsx"
    Workbook().save(workbook)
    content = workbook.read_bytes()
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w") as archive:
        archive.writestr("H040M SO Form.xlsx", content)
        archive.writestr("Q777Q SO Form.xlsx", content)
    files = {
        "order_workbook": [Upload("orders.xlsx", content)],
        "sof_archive": [Upload("sof-files.zip", zip_buffer.getvalue())],
    }

    _, sofs = collect_uploaded_workbooks(files, 20, 100 * 1024 * 1024)

    assert [upload.filename for upload in sofs] == [
        "H040M SO Form.xlsx", "Q777Q SO Form.xlsx",
    ]


def test_worker_rejects_duplicate_sof_filenames(tmp_path):
    workbook = tmp_path / "valid.xlsx"
    Workbook().save(workbook)
    content = workbook.read_bytes()
    files = {
        "order_workbook": [Upload("orders.xlsx", content)],
        "sof_workbook_0": [Upload("H040M SO Form.xlsx", content)],
        "sof_workbook_1": [Upload("h040m so form.xlsx", content)],
    }

    try:
        collect_uploaded_workbooks(files, 20, 100 * 1024 * 1024)
    except AutomationError as exc:
        assert "Duplicate SOF filename" in str(exc)
    else:
        raise AssertionError("Worker accepted duplicate SOF filenames")


def test_n8n_workflow_packages_multiple_sofs_for_worker():
    workflow_path = (
        Path(__file__).resolve().parents[1]
        / "n8n" / "hanger-automation.workflow.json"
    )
    workflow = json.loads(workflow_path.read_text(encoding="utf-8"))
    nodes = {node["name"]: node for node in workflow["nodes"]}
    fields = nodes["Upload Order và SOF"]["parameters"]["formFields"]["values"]
    sof_field = next(field for field in fields if field["fieldName"] == "sof_workbook")
    request_fields = nodes["Gửi File tới Hanger Worker"]["parameters"][
        "bodyParameters"
    ]["parameters"]

    assert sof_field["multipleFiles"] is True
    assert nodes["Đóng gói các SOF"]["parameters"]["binaryPropertyOutput"] == "sof_archive"
    assert any(field.get("name") == "sof_archive" for field in request_fields)


def test_verified_llm_fallback_can_fill_an_unruled_flat_item(tmp_path):
    sof_root = tmp_path / "sof"
    sof_root.mkdir()
    make_hosiery_sof(sof_root)
    order = tmp_path / "orders.xlsx"
    make_order(order, [(711997, "BJ0871", "79", "7-9", "Flat", "")])
    wb = load_workbook(order)
    wb["Orders"].cell(2, HEADERS.index("Product Description") + 1, "6PK CREW SOCK")
    wb.save(order)
    wb.close()
    decision = {
        "status": "MATCHED", "product_category": "OTHER", "sof_hang_flat": "Flat",
        "hanger_code": "NO", "hanger_color": "NO", "color_sizer": "NO",
        "sticker_hanger": "NO", "size_sticker_hanger": "NO",
        "source_sheet": "HOSIERY", "source_cells": "A3:B5", "confidence": 0.96,
        "reasoning": "SOF says hosiery has no hanger and is flat packed.",
    }
    classifier = FakeClassifier(decision)
    rules = Path(__file__).resolve().parents[1] / "rules" / "hanger_rules.json"

    audit = run(order, sof_root, rules, tmp_path / "out", llm_classifier=classifier)

    assert audit["matched"] == 1
    assert audit["row_results"][0]["match_method"] == "LLM_DEEPSEEK_VERIFIED_SOURCE"
    assert audit["row_results"][0]["source_cells"] == "A3:B5"
    result_wb = load_workbook(audit["output_file"])
    columns = {cell.value: cell.column for cell in result_wb["Orders"][1]}
    assert result_wb["Orders"].cell(2, columns["Hanger code"]).value == "NO"
    result_wb.close()


def test_hbe_packaging_sheet_is_included_in_llm_evidence(tmp_path):
    sof = make_hbe_sof(tmp_path)
    wb = load_workbook(sof, data_only=True)
    evidence = _evidence_text(
        wb,
        [{
            "Label": "AK", "Product Description": "FLEECE PANT SET",
            "Size Configuration": "7", "Hang/Flat": "Flat",
        }],
        45_000,
    )
    wb.close()

    assert "[SHEET: PACKAGING]" in evidence
    assert "D24 = AK-KT-KU-KV-KW" in evidence
    assert "E24 = FLATPACKED" in evidence


def test_verified_hbe_flatpack_decision_is_accepted(tmp_path):
    sof = make_hbe_sof(tmp_path)
    wb = load_workbook(sof, data_only=True)
    row = {
        "Account": "H494M", "Label": "AK", "Product Description": "FLEECE PANT SET",
        "Size Configuration": "7", "Hang/Flat": "Flat",
    }
    current = {
        "status": "REVIEW", "validation_note": "Product category is unresolved",
        "product_category": "", "sof_hang_flat": "", "hanger_code": "",
        "hanger_color": "", "color_sizer": "", "sticker_hanger": "",
        "size_sticker_hanger": "", "source_sheet": "", "source_cells": "",
        "match_method": "", "confidence": 0.0,
    }
    decision = {
        "status": "MATCHED", "product_category": "SETS", "sof_hang_flat": "Flat",
        "hanger_code": "NO", "hanger_color": "NO", "color_sizer": "NO",
        "sticker_hanger": "NO", "size_sticker_hanger": "NO",
        "source_sheet": "PACKAGING", "source_cells": "D24:E24",
        "confidence": 0.97, "reasoning": "Label AK is explicitly flatpacked.",
    }

    result = resolve_llm_decision(
        row, current, decision, {"SETS"}, wb, sof, min_confidence=0.85
    )
    wb.close()

    assert result["status"] == "MATCHED"
    assert result["hanger_code"] == "NO"
    assert result["source_sheet"] == "PACKAGING"


def test_hbe_packaging_citation_must_contain_exact_order_label(tmp_path):
    sof = make_hbe_sof(tmp_path)
    wb = load_workbook(sof, data_only=True)
    row = {
        "Account": "H494M", "Label": "G1", "Product Description": "FLEECE PANT SET",
        "Size Configuration": "3T", "Hang/Flat": "Flat",
    }
    current = {
        "status": "REVIEW", "validation_note": "No matching rule",
        "product_category": "", "sof_hang_flat": "", "hanger_code": "",
        "hanger_color": "", "color_sizer": "", "sticker_hanger": "",
        "size_sticker_hanger": "", "source_sheet": "", "source_cells": "",
        "match_method": "", "confidence": 0.0,
    }
    decision = {
        "status": "MATCHED", "product_category": "SETS", "sof_hang_flat": "Flat",
        "hanger_code": "NO", "hanger_color": "NO", "color_sizer": "NO",
        "sticker_hanger": "NO", "size_sticker_hanger": "NO",
        "source_sheet": "PACKAGING", "source_cells": "D24:E24",
        "confidence": 0.99, "reasoning": "Wrong label row.",
    }

    result = resolve_llm_decision(
        row, current, decision, {"SETS"}, wb, sof, min_confidence=0.85
    )
    wb.close()

    assert result["status"] == "REVIEW"
    assert "does not contain order Label G1" in result["validation_note"]


def test_hbe_hanger_code_missing_from_manual_is_rejected(tmp_path):
    sof = make_hbe_sof(tmp_path)
    wb = load_workbook(sof, data_only=True)
    row = {
        "Account": "H494M", "Label": "G1", "Product Description": "FLEECE PANT SET",
        "Size Configuration": "3T", "Hang/Flat": "Hang",
    }
    current = {
        "status": "REVIEW", "validation_note": "No matching rule",
        "product_category": "", "sof_hang_flat": "", "hanger_code": "",
        "hanger_color": "", "color_sizer": "", "sticker_hanger": "",
        "size_sticker_hanger": "", "source_sheet": "", "source_cells": "",
        "match_method": "", "confidence": 0.0,
    }
    decision = {
        "status": "MATCHED", "product_category": "SETS", "sof_hang_flat": "Hang",
        "hanger_code": "6110", "hanger_color": "WHITE", "color_sizer": "WHITE",
        "sticker_hanger": "NO", "size_sticker_hanger": "NO",
        "source_sheet": "PACKAGING", "source_cells": "D25:G25",
        "confidence": 0.99, "reasoning": "External manual is not present.",
    }

    result = resolve_llm_decision(
        row, current, decision, {"SETS"}, wb, sof, min_confidence=0.85
    )
    wb.close()

    assert result["status"] == "REVIEW"
    assert "does not support hanger_code: 6110" in result["validation_note"]


def test_compound_description_resolves_to_its_head_noun():
    allowed = {"BOTTOMS", "TOPS", "SETS", "COVERALLS", "HOSIERY", "OUTERWEAR"}

    # The component garment is named first, the head noun last.
    assert infer_category("KNIT SHORT SET", allowed)[0] == "SETS"
    assert infer_category("FLEECE PANT SET", allowed)[0] == "SETS"
    assert infer_category("8PK CREW SOCK", allowed)[0] == "HOSIERY"
    assert infer_category("BABY KNIT ROMPER", allowed)[0] == "COVERALLS"
    # A plain description still resolves, and an unknown one stays unresolved.
    assert infer_category("KNIT LEGGING", allowed)[0] == "BOTTOMS"
    assert infer_category("MYSTERY ITEM", allowed)[0] == ""


def test_category_is_restricted_to_allowed_sheets():
    assert infer_category("8PK CREW SOCK", {"HOSIERY"})[0] == "HOSIERY"
    # Hosiery is not a top: with HOSIERY disallowed the row stays unresolved
    # rather than being filed under the nearest available category.
    assert infer_category("8PK CREW SOCK", {"TOPS", "SETS"})[0] == ""


def test_label_name_is_read_and_reported(tmp_path):
    sof_root = tmp_path / "sof"
    sof_root.mkdir()
    make_sof(sof_root)
    order = tmp_path / "orders.xlsx"
    make_order(order, [(751419, "26M316", "WH", "2T - 3T - 4T", "Hang", "6110")])
    rules = Path(__file__).resolve().parents[1] / "rules" / "hanger_rules.json"

    audit = run(order, sof_root, rules, tmp_path / "out")

    entry = audit["row_results"][0]
    assert entry["label_name"] == "WITH HANGER"
    wb = load_workbook(Path(audit["result_file"]))
    ws = wb["KET QUA"]
    headers = {c.value: c.column for c in ws[1]}
    assert ws.cell(2, headers["Label Name"]).value == "WITH HANGER"
    wb.close()


def make_split_evidence_sof(root: Path) -> Path:
    """SOF shaped like the real H040M SETS tab.

    The hanger code sits in a table, the colour is stated in a sentence below it,
    and an unrelated "N/A" belongs to a size column that carries no hanger.
    """
    path = root / "H040M Stock Replenishment SO Form 8.31.26.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "SETS"
    ws["A21"] = "Hanger Required"
    ws["B21"] = "Yes"
    ws["A22"] = "Hanger Type for 2 pc pant sets"
    ws["B22"] = "496/9508"
    ws["C22"] = "N/A"
    ws["A23"] = "Sizers"
    ws["B23"] = "White size clip / black lettering"
    ws["B27"] = "Use white plastic hangers on merged size packs, hang the set together."
    wb.save(path)
    return path


def llm_review_state() -> dict:
    return {
        "status": "REVIEW", "validation_note": "No matching rule",
        "product_category": "", "sof_hang_flat": "", "hanger_code": "",
        "hanger_color": "", "color_sizer": "", "sticker_hanger": "",
        "size_sticker_hanger": "", "source_sheet": "", "source_cells": "",
        "match_method": "", "confidence": 0.0,
    }


def test_no_hanger_color_on_a_hang_decision_is_rejected(tmp_path):
    sof = make_split_evidence_sof(tmp_path)
    wb = load_workbook(sof, data_only=True)
    row = {"Account": "H040M", "Label": "WH", "Hang/Flat": "Hang"}
    decision = {
        "status": "MATCHED", "product_category": "SETS", "sof_hang_flat": "Hang",
        "hanger_code": "496/9508", "hanger_color": "NO",
        "color_sizer": "White size clip / black lettering",
        "sticker_hanger": "NO", "size_sticker_hanger": "NO",
        "source_sheet": "SETS", "source_cells": "A21:C23",
        "confidence": 0.99, "reasoning": "Colour is not stated in the table.",
    }

    result = resolve_llm_decision(
        row, llm_review_state(), decision, {"SETS"}, wb, sof, min_confidence=0.85
    )
    wb.close()

    assert result["status"] == "REVIEW"
    assert "Hang rule must state a real hanger_color" in result["validation_note"]
    assert result["hanger_color"] == ""


def test_unrelated_na_does_not_license_a_no_claim(tmp_path):
    sof = make_split_evidence_sof(tmp_path)
    wb = load_workbook(sof, data_only=True)
    row = {"Account": "H040M", "Label": "WH", "Hang/Flat": "Hang"}
    decision = {
        "status": "MATCHED", "product_category": "SETS", "sof_hang_flat": "Hang",
        "hanger_code": "496/9508", "hanger_color": "WHITE", "color_sizer": "NO",
        "sticker_hanger": "NO", "size_sticker_hanger": "NO",
        "source_sheet": "SETS", "source_cells": "A21:C23",
        "confidence": 0.99, "reasoning": "Borrowing the N/A in C22.",
    }

    result = resolve_llm_decision(
        row, llm_review_state(), decision, {"SETS"}, wb, sof, min_confidence=0.85
    )
    wb.close()

    assert result["status"] == "REVIEW"
    assert "does not support color_sizer" in result["validation_note"]


def test_general_white_sentence_cannot_override_the_codes_own_cell(tmp_path):
    """A merged-pack instruction must not recolour a size band the SOF calls black."""
    path = tmp_path / "H040M Stock Replenishment SO Form 8.31.26.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "SETS"
    ws["A21"] = "Hanger Required"
    ws["B21"] = "Yes"
    ws["A26"] = "Hanger Type for 2 pc pant sets"
    ws["B26"] = "472/9510"
    ws["G26"] = "485B/6010B Black Crown Sizer w/ White Lettering"
    ws["A31"] = "Sizers"
    ws["G31"] = "Black size clip w/ white lettering"
    ws["B33"] = "Use white plastic hangers on merged size packs, hang the set."
    wb.save(path)

    wb = load_workbook(path, data_only=True)
    row = {"Account": "H040M", "Label": "WH", "Hang/Flat": "Hang"}
    decision = {
        "status": "MATCHED", "product_category": "SETS", "sof_hang_flat": "Hang",
        "hanger_code": "485B/6010B",
        "color_sizer": "Black size clip w/ white lettering",
        "sticker_hanger": "NO", "size_sticker_hanger": "NO",
        "source_sheet": "SETS", "source_cells": ["A21:G31", "B33:B33"],
        "confidence": 0.99, "reasoning": "Large size band.",
    }

    rejected = resolve_llm_decision(
        row, llm_review_state(), dict(decision, hanger_color="WHITE"),
        {"SETS"}, wb, path, min_confidence=0.85,
    )
    assert rejected["status"] == "REVIEW"
    assert "states a BLACK hanger" in rejected["validation_note"]

    accepted = resolve_llm_decision(
        row, llm_review_state(), dict(decision, hanger_color="BLACK"),
        {"SETS"}, wb, path, min_confidence=0.85,
    )
    wb.close()
    assert accepted["status"] == "MATCHED", accepted["validation_note"]
    assert accepted["hanger_color"] == "BLACK"


def test_not_stated_accessory_means_none_required_but_hanger_is_reviewed(tmp_path):
    """Silence about a sticker is absence; silence about the hanger is a miss."""
    sof = make_split_evidence_sof(tmp_path)
    wb = load_workbook(sof, data_only=True)
    row = {"Account": "H040M", "Label": "WH", "Hang/Flat": "Hang"}
    base = {
        "status": "MATCHED", "product_category": "SETS", "sof_hang_flat": "Hang",
        "hanger_code": "496/9508", "hanger_color": "WHITE",
        "color_sizer": "White size clip / black lettering",
        "source_sheet": "SETS", "source_cells": ["A21:C23", "B27:B27"],
        "confidence": 0.99, "reasoning": "Table plus the sentence below it.",
    }

    accessory = resolve_llm_decision(
        row, llm_review_state(),
        dict(base, sticker_hanger="NOT_STATED", size_sticker_hanger="NOT_STATED"),
        {"SETS"}, wb, sof, min_confidence=0.85,
    )
    assert accessory["status"] == "MATCHED", accessory["validation_note"]
    assert accessory["sticker_hanger"] == "NO"

    hanger = resolve_llm_decision(
        row, llm_review_state(),
        dict(base, hanger_color="NOT_STATED", sticker_hanger="NO",
             size_sticker_hanger="NO"),
        {"SETS"}, wb, sof, min_confidence=0.85,
    )
    wb.close()
    assert hanger["status"] == "REVIEW"
    assert "SOF does not state: hanger_color" in hanger["validation_note"]


def test_a_decision_may_cite_several_ranges(tmp_path):
    sof = make_split_evidence_sof(tmp_path)
    wb = load_workbook(sof, data_only=True)
    row = {"Account": "H040M", "Label": "WH", "Hang/Flat": "Hang"}
    decision = {
        "status": "MATCHED", "product_category": "SETS", "sof_hang_flat": "Hang",
        "hanger_code": "496/9508", "hanger_color": "WHITE",
        "color_sizer": "White size clip / black lettering",
        "sticker_hanger": "NO", "size_sticker_hanger": "NO",
        "source_sheet": "SETS", "source_cells": ["A21:C23", "B27:B27"],
        "confidence": 0.99, "reasoning": "Table plus the colour sentence below it.",
    }

    result = resolve_llm_decision(
        row, llm_review_state(), decision, {"SETS"}, wb, sof, min_confidence=0.85
    )
    wb.close()

    assert result["status"] == "MATCHED", result["validation_note"]
    assert result["hanger_color"] == "WHITE"
    assert result["source_cells"] == "A21:C23,B27:B27"


def test_llm_hallucinated_value_is_rejected(tmp_path):
    sof_root = tmp_path / "sof"
    sof_root.mkdir()
    make_sof(sof_root)
    order = tmp_path / "orders.xlsx"
    make_order(order, [(711999, "UNKNOWN", "WH", "2T - 3T", "Hang", "")])
    decision = {
        "status": "MATCHED", "product_category": "BOTTOMS", "sof_hang_flat": "Hang",
        "hanger_code": "9999", "hanger_color": "BLUE", "color_sizer": "Blue clip",
        "sticker_hanger": "NO", "size_sticker_hanger": "NO",
        "source_sheet": "BOTTOMS", "source_cells": "B15:G18", "confidence": 0.99,
        "reasoning": "Unsupported values.",
    }
    rules = Path(__file__).resolve().parents[1] / "rules" / "hanger_rules.json"

    audit = run(
        order, sof_root, rules, tmp_path / "out", llm_classifier=FakeClassifier(decision)
    )

    assert audit["review"] == 1
    assert "Rejected DeepSeek decision" in audit["row_results"][0]["validation_note"]
    assert "hanger_code" in audit["row_results"][0]["validation_note"]


def capture_request_payload(text_mode: bool = False) -> dict:
    """Build one real API payload without letting it reach the network."""
    import urllib.request

    classifier = DeepSeekClassifier(DeepSeekSettings(api_keys=("test",)))
    captured = {}

    def fake_urlopen(request, timeout=None):
        captured["payload"] = json.loads(request.data.decode("utf-8"))
        raise RuntimeError("stop before the network")

    groups = {"G0001": {"row_numbers": [3], "row": {
        "Account": "H040M", "Label": "-V", "Label Name": "ST HANG ALT PACK",
        "Product Description": "KNIT SHORT SET", "Size Configuration": "12M",
        "Hang/Flat": "Hang",
    }}}
    evidence = "\n[SHEET: SETS]\nA22 = Hanger Type\nB22 = 496/9508\n"
    original = urllib.request.urlopen
    urllib.request.urlopen = fake_urlopen
    try:
        classifier._request(
            groups, {"SETS"}, "H040M.xlsx", evidence, "test", text_mode=text_mode
        )
    except RuntimeError:
        pass
    finally:
        urllib.request.urlopen = original
    return captured["payload"]


def test_schema_asks_for_the_citation_before_the_values():
    """Generation is left to right, so the model must cite before it commits."""
    payload = capture_request_payload()
    fields = list(json.loads(payload["messages"][1]["content"])
                  ["output_schema"]["decisions"][0])

    assert fields.index("reasoning") < fields.index("hanger_code")
    assert fields.index("source_cells") < fields.index("hanger_code")
    assert fields.index("hanger_code") < fields.index("status")


def test_text_mode_swaps_the_citation_fields_but_keeps_the_order():
    payload = capture_request_payload(text_mode=True)
    fields = list(json.loads(payload["messages"][1]["content"])
                  ["output_schema"]["decisions"][0])

    assert "source_cells" not in fields and "source_sheet" not in fields
    assert fields.index("source_quote") < fields.index("hanger_code")


def test_label_name_is_sent_to_the_model():
    payload = capture_request_payload()
    group = json.loads(payload["messages"][1]["content"])["order_groups"][0]

    assert group["label_name"] == "ST HANG ALT PACK"


def test_deepseek_groups_duplicate_order_patterns_before_api_call(tmp_path):
    sof = make_hosiery_sof(tmp_path)
    wb = load_workbook(sof, data_only=True)
    settings = DeepSeekSettings(api_keys=("test",), batch_size=12, max_calls=2)
    classifier = DeepSeekClassifier(settings)
    captured = []

    def fake_request(groups, allowed, source_name, evidence, api_key):
        captured.append(groups)
        return {"decisions": [
            {"group_id": group_id, "status": "REVIEW", "reasoning": "test"}
            for group_id in groups
        ]}

    classifier._request = fake_request
    row = {
        "Account": "H040M", "Division": "G", "Label": "79", "Ref#": "R1",
        "Style": "S1", "Product Description": "6PK CREW SOCK",
        "Size Configuration": "7-9", "Hang/Flat": "Flat",
    }
    decisions = classifier.classify([(2, row), (3, dict(row))], {"OTHER"}, wb, sof)
    wb.close()

    assert len(captured) == 1
    assert len(captured[0]) == 1
    assert set(decisions) == {2, 3}
    assert classifier.groups_requested == 1


def test_unique_pending_sheet_is_preferred_over_completed_history(tmp_path):
    sof_root = tmp_path / "sof"
    sof_root.mkdir()
    make_sof(sof_root)
    order = tmp_path / "multi-sheet.xlsx"
    wb = Workbook()
    completed = wb.active
    completed.title = "Completed"
    completed.append(HEADERS)
    completed.append(["" for _ in HEADERS])
    pending = wb.create_sheet("Pending")
    pending_headers = HEADERS[:HEADERS.index("Hanger code")]
    pending.append(pending_headers)
    values = {
        "Division": "G", "Ref#": "NKG-LEM316", "Style": "26M316",
        "Product Description": "KNIT LEGGING", "Label": "WH", "Account": "H040M",
        "PO#": 751419, "Size Configuration": "2T - 3T - 4T", "Hang/Flat": "Hang",
    }
    pending.append([values.get(header, "") for header in pending_headers])
    wb.save(order)
    wb.close()
    rules = Path(__file__).resolve().parents[1] / "rules" / "hanger_rules.json"

    audit = run(order, sof_root, rules, tmp_path / "out")

    assert audit["matched"] == 1
    result_wb = load_workbook(audit["output_file"])
    assert result_wb["HANGER REVIEW"].max_row == 1
    pending_columns = {cell.value: cell.column for cell in result_wb["Pending"][1]}
    assert result_wb["Pending"].cell(2, pending_columns["Hanger code"]).value == "6110"
    result_wb.close()


def test_multiple_api_keys_run_batches_in_parallel(tmp_path):
    sof = make_hosiery_sof(tmp_path)
    wb = load_workbook(sof, data_only=True)
    settings = DeepSeekSettings(
        api_keys=("key-a", "key-b"), batch_size=1, max_calls=10
    )
    classifier = DeepSeekClassifier(settings)
    barrier = threading.Barrier(2, timeout=5)
    used_keys = []
    lock = threading.Lock()

    def fake_request(groups, allowed, source_name, evidence, api_key):
        with lock:
            used_keys.append(api_key)
        # Deadlocks unless two batches are genuinely in flight together.
        barrier.wait()
        return {"decisions": [
            {"group_id": group_id, "status": "REVIEW", "reasoning": "test"}
            for group_id in groups
        ]}

    classifier._request = fake_request
    rows = []
    for index in range(2):
        rows.append((index + 2, {
            "Account": "H040M", "Division": "G", "Label": "79", "Ref#": f"R{index}",
            "Style": f"S{index}", "Product Description": "6PK CREW SOCK",
            "Color": "023", "Size Configuration": "", "Hang/Flat": "",
        }))

    classifier.classify(rows, {"HOSIERY"}, wb, Path(sof))
    wb.close()

    assert sorted(used_keys) == ["key-a", "key-b"]


def test_api_keys_parsed_from_comma_list_and_numbered_slots(monkeypatch):
    monkeypatch.setenv("HANGER_LLM_ENABLED", "true")
    monkeypatch.setenv("DEEPSEEK_API_KEY", " key-one , key-two ")
    monkeypatch.setenv("DEEPSEEK_API_KEY_2", "key-three")
    monkeypatch.setenv("DEEPSEEK_API_KEY_3", "key-two")  # duplicate is dropped
    settings = DeepSeekSettings.from_env()
    assert settings.api_keys == ("key-one", "key-two", "key-three")
    assert settings.min_confidence == 0.65


def test_deepseek_http_402_is_not_retried(monkeypatch):
    import urllib.error
    import urllib.request

    attempts = 0

    def no_credit(request, timeout=None):
        nonlocal attempts
        attempts += 1
        raise urllib.error.HTTPError(
            request.full_url, 402, "Payment Required", {}, None
        )

    monkeypatch.setattr(urllib.request, "urlopen", no_credit)
    classifier = DeepSeekClassifier(
        DeepSeekSettings(api_keys=("test",), max_calls=5)
    )
    groups = {"G0001": {"row_numbers": [2], "row": {
        "Account": "H040M", "Label": "WH", "Product Description": "TEE",
    }}}

    try:
        classifier._request(groups, {"TOPS"}, "H040M.xlsx", "", "test")
    except RuntimeError as exc:
        assert str(exc) == "DeepSeek HTTP 402"
    else:
        raise AssertionError("HTTP 402 should fail the batch")

    assert attempts == 1
    assert classifier.calls == 1


def test_sof_sourced_columns_get_yellow_header(tmp_path):
    sof_root = tmp_path / "sof"
    sof_root.mkdir()
    make_sof(sof_root)
    order = tmp_path / "3. TONG HOP HANGER-2026.xlsx"
    make_order(order)
    rules = Path(__file__).resolve().parents[1] / "rules" / "hanger_rules.json"

    audit = run(order, sof_root, rules, tmp_path / "out")

    result_wb = load_workbook(audit["output_file"])
    ws = result_wb[result_wb.sheetnames[0]]
    header_row = 1
    headers = {cell.value: cell.column for cell in ws[header_row]}
    for field in ("Hanger code", "Hanger color", "Color Sizer",
                  "Sticker hanger", "SIZE Sticker hanger"):
        cell = ws.cell(header_row, headers[field])
        assert cell.fill.fill_type == "solid"
        assert cell.fill.fgColor.rgb.endswith("FFFF00"), field
    result_wb.close()


class FakeRedis:
    """Minimal stand-in for the redis client surface DecisionCache uses."""

    def __init__(self, broken=False):
        self.store = {}
        self.broken = broken
        self.writes = 0

    def mget(self, keys):
        if self.broken:
            raise ConnectionError("redis down")
        return [self.store.get(k) for k in keys]

    def pipeline(self, transaction=False):
        return FakePipeline(self)


class FakePipeline:
    def __init__(self, parent):
        self.parent = parent
        self.queued = []

    def set(self, key, value, ex=None):
        self.queued.append((key, value))

    def execute(self):
        if self.parent.broken:
            raise ConnectionError("redis down")
        for key, value in self.queued:
            self.parent.store[key] = value
            self.parent.writes += 1


def _classifier_rows():
    return [(2, {
        "Account": "H040M", "Division": "G", "Label": "79", "Ref#": "R1",
        "Style": "S1", "Product Description": "6PK CREW SOCK",
        "Color": "023", "Size Configuration": "", "Hang/Flat": "",
    })]


def test_cached_decision_is_reused_without_calling_the_api(tmp_path):
    sof = make_hosiery_sof(tmp_path)
    wb = load_workbook(sof, data_only=True)
    fake = FakeRedis()
    settings = DeepSeekSettings(api_keys=("k",), batch_size=12, max_calls=5)

    calls = []

    def fake_request(groups, allowed, source_name, evidence, api_key):
        calls.append(groups)
        return {"decisions": [
            {"group_id": gid, "status": "MATCHED", "hanger_code": "6110"}
            for gid in groups
        ]}

    first = DeepSeekClassifier(settings, DecisionCache(fake, 3600))
    first._request = fake_request
    first.classify(_classifier_rows(), {"HOSIERY"}, wb, Path(sof))
    assert len(calls) == 1 and fake.writes == 1

    # A second classifier sharing the cache must not hit the API at all.
    second = DeepSeekClassifier(settings, DecisionCache(fake, 3600))
    second._request = fake_request
    out = second.classify(_classifier_rows(), {"HOSIERY"}, wb, Path(sof))
    wb.close()

    assert len(calls) == 1, "cached group was re-sent to the API"
    assert out[2]["hanger_code"] == "6110"
    assert second.cache.hits == 1


def test_changed_sof_contents_do_not_reuse_cached_answers(tmp_path):
    sof_a = make_hosiery_sof(tmp_path)
    wb = load_workbook(sof_a, data_only=True)
    fake = FakeRedis()
    settings = DeepSeekSettings(api_keys=("k",), batch_size=12, max_calls=5)
    key_a = DecisionCache.build_key("m", "digestA", ("x",), {"HOSIERY"})
    key_b = DecisionCache.build_key("m", "digestB", ("x",), {"HOSIERY"})
    assert key_a != key_b

    other = tmp_path / "other.xlsx"
    other.write_bytes(sof_a.read_bytes() + b"different")
    from deepseek_classifier import _file_digest
    assert _file_digest(sof_a) != _file_digest(other)
    wb.close()


def test_broken_cache_degrades_to_plain_api_calls(tmp_path):
    sof = make_hosiery_sof(tmp_path)
    wb = load_workbook(sof, data_only=True)
    fake = FakeRedis(broken=True)
    settings = DeepSeekSettings(api_keys=("k",), batch_size=12, max_calls=5)
    classifier = DeepSeekClassifier(settings, DecisionCache(fake, 3600))
    calls = []

    def fake_request(groups, allowed, source_name, evidence, api_key):
        calls.append(groups)
        return {"decisions": [
            {"group_id": gid, "status": "MATCHED", "hanger_code": "6110"}
            for gid in groups
        ]}

    classifier._request = fake_request
    out = classifier.classify(_classifier_rows(), {"HOSIERY"}, wb, Path(sof))
    wb.close()

    assert len(calls) == 1, "a dead cache must not stop the API path"
    assert out[2]["hanger_code"] == "6110"
    assert classifier.cache.errors >= 1


def test_result_sheet_has_tong_hop_columns_and_header_colours(tmp_path):
    sof_root = tmp_path / "sof"
    sof_root.mkdir()
    make_sof(sof_root)
    order = tmp_path / "3. TONG HOP HANGER-2026.xlsx"
    make_order(order)
    rules = Path(__file__).resolve().parents[1] / "rules" / "hanger_rules.json"

    audit = run(order, sof_root, rules, tmp_path / "out")

    result = Path(audit["result_file"])
    assert result.is_file()
    wb = load_workbook(result)
    assert wb.sheetnames == ["KET QUA"], "deliverable must be a single tab"
    ws = wb["KET QUA"]
    headers = [c.value for c in ws[1]]
    assert headers == [
        "Division", "Season", "Year", "Ref#", "Style", "Product Description",
        "Color", "Label", "Label Name", "Account", "PO#", "PO Qty",
        "Size Configuration", "Pack Ratio", "Master Box Quantity", "Hang/Flat",
        "Hanger code", "Hanger color", "Color Sizer", "Sticker hanger",
        "SIZE Sticker hanger", "NCC",
        "Status", "Validation Note", "Match Method", "Confidence",
        "SOF File", "SOF Sheet", "SOF Cells", "SOF Section", "SOF Lines", "SOF Excerpt",
    ]
    for cell in ws[1]:
        if cell.column <= 16:
            expected = "33CCCC"
        elif cell.column <= 22:
            expected = "FFFF00"
        else:
            expected = "D9D9D9"
        assert cell.fill.fill_type == "solid"
        assert cell.fill.fgColor.rgb.endswith(expected), cell.value
    assert ws.max_row == len(audit["row_results"]) + 1
    wb.close()


def test_result_sheet_leaves_unverified_rows_blank(tmp_path):
    sof_root = tmp_path / "sof"
    sof_root.mkdir()
    make_sof(sof_root)
    order = tmp_path / "3. TONG HOP HANGER-2026.xlsx"
    make_order(order)
    rules = Path(__file__).resolve().parents[1] / "rules" / "hanger_rules.json"

    audit = run(order, sof_root, rules, tmp_path / "out")
    wb = load_workbook(Path(audit["result_file"]))
    ws = wb["KET QUA"]
    headers = {c.value: c.column for c in ws[1]}
    for offset, entry in enumerate(audit["row_results"], start=2):
        value = ws.cell(offset, headers["Hanger code"]).value or ""
        if entry["status"] == "MATCHED":
            assert value == entry["hanger_code"]
        else:
            assert value == "", "unverified row must not carry a hanger value"
        assert (ws.cell(offset, headers["NCC"]).value or "") == ""
    wb.close()


def test_result_sheet_explains_every_row_including_reviews(tmp_path):
    sof_root = tmp_path / "sof"
    sof_root.mkdir()
    make_sof(sof_root)
    order = tmp_path / "3. TONG HOP HANGER-2026.xlsx"
    make_order(order)
    rules = Path(__file__).resolve().parents[1] / "rules" / "hanger_rules.json"

    audit = run(order, sof_root, rules, tmp_path / "out")
    wb = load_workbook(Path(audit["result_file"]))
    ws = wb["KET QUA"]
    headers = {c.value: c.column for c in ws[1]}

    for offset, entry in enumerate(audit["row_results"], start=2):
        # Status and its explanation are present on every row, matched or not.
        assert ws.cell(offset, headers["Status"]).value == entry["status"]
        assert (ws.cell(offset, headers["Validation Note"]).value or "") == entry["validation_note"]
        assert (ws.cell(offset, headers["SOF Cells"]).value or "") == entry["source_cells"]
    wb.close()


def test_rows_of_other_accounts_are_out_of_scope_not_review(tmp_path):
    """Uploading one account's SOF must check that account only."""
    sof_root = tmp_path / "sof"
    sof_root.mkdir()
    make_sof(sof_root)  # H040M only
    order = tmp_path / "3. TONG HOP HANGER-2026.xlsx"
    make_order(order)

    # Add rows for an account whose SOF was not uploaded.
    wb = load_workbook(order)
    ws = wb[wb.sheetnames[0]]
    headers = {c.value: c.column for c in ws[1]}
    start = ws.max_row + 1
    for offset in range(3):
        target = start + offset
        ws.cell(target, headers["Account"], "Z999M")
        ws.cell(target, headers["PO#"], "0999999")
        ws.cell(target, headers["Ref#"], "ZZZ-1")
        ws.cell(target, headers["Style"], "ZS1")
        ws.cell(target, headers["Product Description"], "KNIT LEGGING")
    wb.save(order)
    wb.close()

    rules = Path(__file__).resolve().parents[1] / "rules" / "hanger_rules.json"
    audit = run(order, sof_root, rules, tmp_path / "out")

    assert audit["accounts_checked"] == ["H040M"]
    assert "Z999M" in audit["accounts_skipped"]
    assert audit["skipped_no_sof"] == 3
    accounts_in_results = {r["account"].upper() for r in audit["row_results"]}
    assert "Z999M" not in accounts_in_results, "out-of-scope rows must not be reported"
    assert all(r["account"].upper() == "H040M" for r in audit["row_results"])
    # The skipped rows are still recorded, so nothing is silently dropped.
    assert sum(1 for r in audit["skipped_rows"] if r["account"] == "Z999M") == 3

    # The deliverable carries only the account that was actually checked.
    result_wb = load_workbook(Path(audit["result_file"]))
    ws = result_wb["KET QUA"]
    headers = {c.value: c.column for c in ws[1]}
    seen = {ws.cell(r, headers["Account"]).value for r in range(2, ws.max_row + 1)}
    assert seen == {"H040M"}
    result_wb.close()


def test_run_reports_an_error_when_no_account_matches_the_uploaded_sof(tmp_path):
    sof_root = tmp_path / "sof"
    sof_root.mkdir()
    make_sof(sof_root, name="Q777Q Stock Replenishment SO Form 8.31.26.xlsx")
    order = tmp_path / "3. TONG HOP HANGER-2026.xlsx"
    make_order(order)
    rules = Path(__file__).resolve().parents[1] / "rules" / "hanger_rules.json"

    audit = run(order, sof_root, rules, tmp_path / "out")

    assert audit["accounts_checked"] == []
    assert audit["row_results"] == []
    assert any("nothing was checked" in e for e in audit["errors"])


def test_general_packing_sheet_is_not_accepted_as_a_citation(tmp_path):
    """The REPLESNISHMENT regression: a general rule is not per-style evidence."""
    sof_root = tmp_path / "sof"
    sof_root.mkdir()
    sof = make_sof(sof_root)

    # Give the SOF a general packing sheet that says Flat, like the real one.
    wb = load_workbook(sof)
    ws = wb.create_sheet("REPLESNISHMENT")
    ws["A17"] = "Division G (NIKE)"
    ws["A18"] = "Apparel"
    ws["B18"] = ("Flat Pack, individual sealed polybag with hangtag/UPC visible "
                 "through polybag except the categories below.")
    wb.save(sof)
    wb.close()

    order = tmp_path / "orders.xlsx"
    # A style with no approved rule, so the row reaches the LLM path.
    make_order(order, [(711999, "UNKNOWN", "WH", "2T - 3T", "Hang", "")])
    rules = Path(__file__).resolve().parents[1] / "rules" / "hanger_rules.json"

    decision = {
        "status": "MATCHED", "product_category": "BOTTOMS", "sof_hang_flat": "Flat",
        "hanger_code": "NO", "hanger_color": "NO", "color_sizer": "NO",
        "sticker_hanger": "NO", "size_sticker_hanger": "NO",
        "source_sheet": "REPLESNISHMENT", "source_cells": "A17:B18",
        "confidence": 0.95, "reasoning": "Flat Pack stated for Division G apparel.",
    }
    audit = run(order, sof_root, rules, tmp_path / "out",
                llm_classifier=FakeClassifier(decision))

    assert audit["mismatch"] == 0, "a general packing rule must not create a MISMATCH"
    notes = " ".join(r["validation_note"] for r in audit["row_results"])
    assert "general packing rules" in notes
    assert all(r["match_method"] != "LLM_DEEPSEEK_VERIFIED_SOURCE"
               for r in audit["row_results"])


def test_category_sheet_citation_is_still_accepted(tmp_path):
    """The gate must reject only non-category sheets, not every citation."""
    sof_root = tmp_path / "sof"
    sof_root.mkdir()
    make_sof(sof_root)
    order = tmp_path / "3. TONG HOP HANGER-2026.xlsx"
    make_order(order)
    rules = Path(__file__).resolve().parents[1] / "rules" / "hanger_rules.json"

    from hanger_automation import CATEGORY_SOURCE_SHEETS
    assert "BOTTOMS" in CATEGORY_SOURCE_SHEETS
    assert "REPLESNISHMENT" not in CATEGORY_SOURCE_SHEETS
    assert "GENERAL INFO" not in CATEGORY_SOURCE_SHEETS

    audit = run(order, sof_root, rules, tmp_path / "out")
    # The approved rules cite BOTTOMS and must keep working.
    assert audit["matched"] > 0
